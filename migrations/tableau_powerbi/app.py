import re
import json
import zipfile
import difflib
from io import BytesIO
from html import escape
from pathlib import Path
from collections import Counter, defaultdict
from typing import Literal, Optional
from xml.etree import ElementTree as ET

import pandas as pd
import streamlit as st
try:
    from openai import OpenAI
except Exception:
    OpenAI = None
from pydantic import BaseModel, Field, ValidationError
from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter


# ===========================================================================
# STYLING CONSTANTS (Excel)
# ===========================================================================

THIN = Side(style='thin', color='CCCCCC')
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WRAP_TOP = Alignment(wrap_text=True, vertical='top')

COLORS = {
    'header': ('1F4E79', 'FFFFFF'),
    'table': ('DDEBF7', '000000'),
    'calc_measure': ('E2EFDA', '000000'),
    'calc_lod': ('FFF2CC', '000000'),
    'filter_shared': ('EBF3FB', '000000'),
    'filter_worksheet': ('F4ECFA', '000000'),
    'grp_a': ('FFE699', '7B5B00'),
    'grp_b': ('C6EFCE', '276221'),
    'grp_c': ('DAEEF3', '17375E'),
    'unique': ('F2F2F2', '7F7F7F'),
    'unparsed': ('EDEDED', 'AAAAAA'),
}

FAST_DATA_FORMATTING = True
EXCEL_CELL_LIMIT = 32767

TABLES_HEADERS = ['File Path', 'Workbook Name', 'Datasource', 'Table', 'Connection Class']
TABLES_WIDTHS = [42, 40, 26, 32, 20]

RELATIONSHIPS_HEADERS = ['File Path', 'Workbook Name', 'Left Field', 'Operator', 'Right Field', 'Style']
RELATIONSHIPS_WIDTHS = [42, 40, 28, 12, 28, 30]

FORMULA_HEADERS = [
    'File Path', 'Workbook Name', 'Field', 'Formula Type', 'Raw Formula',
    'Resolved Formula', 'References', 'Power BI Guidance',
]
FORMULA_WIDTHS = [42, 40, 26, 26, 55, 55, 30, 60]

LOD_HEADERS = [
    'File Path', 'Workbook Name', 'Field', 'LOD Type', 'Dimensions',
    'Tableau Expression', 'Suggested DAX', 'Caveat',
]
LOD_WIDTHS = [42, 40, 22, 14, 26, 40, 55, 70]

FILTERS_HEADERS = ['File Path', 'Workbook Name', 'Scope', 'Owner', 'Field', 'Function', 'Power BI Guidance']
FILTERS_WIDTHS = [42, 40, 20, 28, 24, 20, 55]

WORKSHEETS_HEADERS = ['File Path', 'Workbook Name', 'Worksheet', 'Mark Type', 'Suggested Power BI Visual', 'Fields Used']
WORKSHEETS_WIDTHS = [42, 40, 28, 16, 40, 60]

DASHBOARDS_HEADERS = ['File Path', 'Workbook Name', 'Dashboard', 'Worksheets Placed']
DASHBOARDS_WIDTHS = [42, 40, 28, 70]

ACTIONS_HEADERS = ['File Path', 'Workbook Name', 'Action', 'Kind', 'Source Worksheet', 'Activation', 'Power BI Equivalent']
ACTIONS_WIDTHS = [42, 40, 24, 26, 24, 16, 55]

DUPLICATE_HEADERS = [
    'Duplicate Type', 'Group ID', 'File Path', 'Workbook Name',
    'Similarity', 'Matched With', 'Difference Summary',
]
DUPLICATE_WIDTHS = [18, 18, 42, 40, 12, 40, 55]

USAGE_HEADERS = ['Content Usage Entry', 'Matched Workbook File', 'Match Status']
USAGE_WIDTHS = [55, 55, 16]

TABLE_DB_MAPPING_HEADERS = [
    'Unique Table from KashMap', 'Actual Database Table', 'Connection',
    'TDS/TDSX File', 'Evidence / Notes',
]
TABLE_DB_MAPPING_WIDTHS = [40, 40, 20, 45, 70]

VALIDATION_HEADERS = ['Workbook Name', 'Check Category', 'Validation Check', 'Trigger Rule']
VALIDATION_WIDTHS = [30, 26, 90, 36]

FINAL_SUMMARY_HEADERS = ['Metric', 'Count']
FINAL_SUMMARY_WIDTHS = [48, 38]


# ===========================================================================
# AI BUILD PLAN — Pydantic schema, validation-against-facts, prompt
# ===========================================================================

Classification = Literal[
    "confirmed", "rule_based", "inferred", "suggested", "manual_review",
]

PowerBILayer = Literal[
    "Source SQL", "Power Query", "Mapping Table", "Semantic Model",
    "DAX Measure", "Calculated Column", "Report Filter", "Slicer",
    "Visual", "Drill-through", "Paginated Report", "Manual Review",
]


class PlanStep(BaseModel):
    step_number: int = Field(ge=1)
    title: str
    action: str
    power_bi_layer: PowerBILayer
    classification: Classification
    reason: str
    referenced_sources: list[str] = Field(default_factory=list)
    referenced_fields: list[str] = Field(default_factory=list)
    confidence: int = Field(ge=0, le=100)
    manual_validation: Optional[str] = None


class ReportSummary(BaseModel):
    report_name: str
    likely_purpose: str
    report_type: str
    complexity_level: Literal["Very Small", "Small", "Medium", "Large", "Very Complex"]
    metadata_level: Literal["twb_only"] = "twb_only"
    overall_confidence: int = Field(ge=0, le=100)


class BuildPlan(BaseModel):
    report_summary: ReportSummary
    implementation_plan: list[PlanStep]
    data_preparation: list[PlanStep] = Field(default_factory=list)
    calculation_plan: list[PlanStep] = Field(default_factory=list)
    filter_parameter_plan: list[PlanStep] = Field(default_factory=list)
    page_recommendations: list[PlanStep] = Field(default_factory=list)
    validation_checks: list[str] = Field(default_factory=list)
    manual_review_items: list[str] = Field(default_factory=list)
    limitations: list[str] = Field(default_factory=list)


def validate_plan_references(plan: BuildPlan, allowed_sources: set, allowed_fields: set) -> list[str]:
    """Cross-checks every table/field the AI referenced against the actual
    parsed facts for this workbook. Anything not present gets flagged so it
    can be stripped/downgraded rather than silently trusted."""
    errors: list[str] = []
    sections = [
        plan.implementation_plan, plan.data_preparation, plan.calculation_plan,
        plan.filter_parameter_plan, plan.page_recommendations,
    ]
    for section in sections:
        for step in section:
            unknown_sources = set(step.referenced_sources) - allowed_sources
            unknown_fields = set(step.referenced_fields) - allowed_fields
            if unknown_sources:
                errors.append(f"Step {step.step_number} references unknown sources: {sorted(unknown_sources)}")
            if unknown_fields:
                errors.append(f"Step {step.step_number} references unknown fields: {sorted(unknown_fields)}")
    return errors


def strip_unverified_steps(plan: BuildPlan, allowed_sources: set, allowed_fields: set) -> BuildPlan:
    """Downgrades any step referencing unknown sources/fields to manual_review
    classification and appends an explanatory note, instead of deleting it
    outright (keeps the plan complete while flagging what needs a human)."""
    def clean_section(section: list[PlanStep]) -> list[PlanStep]:
        cleaned = []
        for step in section:
            unknown_sources = set(step.referenced_sources) - allowed_sources
            unknown_fields = set(step.referenced_fields) - allowed_fields
            if unknown_sources or unknown_fields:
                step.classification = "manual_review"
                note = "AI recommendation flagged: referenced "
                parts = []
                if unknown_sources:
                    parts.append(f"unknown source(s) {sorted(unknown_sources)}")
                if unknown_fields:
                    parts.append(f"unknown field(s) {sorted(unknown_fields)}")
                step.manual_validation = note + " and ".join(parts) + " not present in the parsed workbook."
                step.confidence = min(step.confidence, 40)
            cleaned.append(step)
        return cleaned

    plan.implementation_plan = clean_section(plan.implementation_plan)
    plan.data_preparation = clean_section(plan.data_preparation)
    plan.calculation_plan = clean_section(plan.calculation_plan)
    plan.filter_parameter_plan = clean_section(plan.filter_parameter_plan)
    plan.page_recommendations = clean_section(plan.page_recommendations)
    return plan


VALID_POWER_BI_LAYERS = set(PowerBILayer.__args__)
_LAYER_PRECEDENCE = list(PowerBILayer.__args__)


def _coerce_power_bi_layer(value):
    if value in VALID_POWER_BI_LAYERS:
        return value
    parts = re.split(r'[|/,]', str(value))
    parts = [p.strip() for p in parts if p.strip() in VALID_POWER_BI_LAYERS]
    if parts:
        for preferred in _LAYER_PRECEDENCE:
            if preferred in parts:
                return preferred
        return parts[0]
    return "Manual Review"


def normalize_power_bi_layers(raw):
    section_keys = [
        'implementation_plan', 'data_preparation', 'calculation_plan',
        'filter_parameter_plan', 'page_recommendations',
    ]
    for key in section_keys:
        steps = raw.get(key)
        if not isinstance(steps, list):
            continue
        for step in steps:
            if isinstance(step, dict) and 'power_bi_layer' in step:
                step['power_bi_layer'] = _coerce_power_bi_layer(step['power_bi_layer'])
    return raw


AI_BUILD_PLAN_SYSTEM_PROMPT = """You are a senior BI migration engineer converting one parsed Tableau workbook to Power BI.
The input contains .twb-derived workbook logic only (no verified database primary keys, cardinality, table grain,
or complete relationships). rule_findings and lod_translations, if present in the input, are deterministic and
authoritative - never re-derive or contradict them; they were already computed correctly by a rule engine.

Important Tableau-specific facts to reason with:
- Tableau's modern "relationships" (noodle joins) do NOT declare cardinality or join type the way SQL joins do.
  Every relationship in the input should be treated as needing cardinality/uniqueness verification before it
  becomes a Power BI model relationship - flag this explicitly, do not assume one-to-many.
- Table calculations (WINDOW_/RUNNING_/RANK/etc.) depend on the partitioning/addressing of the ORIGINAL VIEW,
  which is not fully recoverable from the XML alone - always classify these manual_review, never confirmed.
- A Measure Names/Values filter is not a row filter - it controls which measures are swappable in a view and
  should map to a Power BI Field Parameter, not a slicer on a real column.

Rules:
1. Use only tables, sources, fields, and calculations present in the input. Never invent a table, field,
   measure, relationship, or business rule.
2. Do not claim primary keys, cardinality, fact/dimension roles, or table grain unless explicitly supplied.
3. Label every implementation step as one of: confirmed, rule_based, inferred, suggested, manual_review.
4. Every step must include referenced_sources and referenced_fields drawn only from the input - this is used
   for automated validation, so accuracy here is critical.
5. Any uncertainty must go into manual_review_items, not be silently guessed.
6. Confidence (0-100) must reflect how directly the input supports the recommendation, not general optimism.
7. Return ONLY strict JSON matching the required schema - no markdown fences, no commentary outside the JSON.

Return JSON with exactly these keys:
{
  "report_summary": {"report_name": "", "likely_purpose": "", "report_type": "",
                      "complexity_level": "Very Small|Small|Medium|Large|Very Complex",
                      "metadata_level": "twb_only", "overall_confidence": 0},
  "implementation_plan": [{"step_number": 1, "title": "", "action": "",
                            "power_bi_layer": "Source SQL|Power Query|Mapping Table|Semantic Model|DAX Measure|Calculated Column|Report Filter|Slicer|Visual|Drill-through|Paginated Report|Manual Review",
                            "classification": "confirmed|rule_based|inferred|suggested|manual_review",
                            "reason": "", "referenced_sources": [], "referenced_fields": [],
                            "confidence": 0, "manual_validation": null}],
  "data_preparation": [... same PlanStep shape ...],
  "calculation_plan": [... same PlanStep shape ..., must cover every item in unresolved_formulas],
  "filter_parameter_plan": [... same PlanStep shape ..., must cover flagged_filters and relationships_needing_cardinality_check],
  "page_recommendations": [... same PlanStep shape ..., cover worksheets/dashboards],
  "validation_checks": ["..."],
  "manual_review_items": ["..."],
  "limitations": ["..."]
}

Every entry in relationships_needing_cardinality_check must appear in filter_parameter_plan or
implementation_plan with classification "manual_review" and a reason explaining that Tableau relationships
don't declare cardinality. Every entry in unresolved_formulas must get exactly one calculation_plan step.
Every lod_translations entry must be referenced (by field name) in calculation_plan with classification
"rule_based" and confidence reflecting the caveat provided (lower confidence when a caveat is present).
"""

_COMPLEXITY_ALIASES = {'High': 'Very Complex', 'Low': 'Small'}


def get_openai_client():
    if OpenAI is None:
        return None
    api_key = st.secrets.get('OPENAI_API_KEY', '')
    if not api_key or api_key == 'paste_your_openai_api_key_here':
        return None
    return OpenAI(api_key=api_key)


def parse_llm_json(text):
    text = (text or '').strip()
    if not text:
        return {}
    try:
        return json.loads(text)
    except json.JSONDecodeError:
        pass
    match = re.search(r'\{.*\}', text, flags=re.DOTALL)
    if match:
        try:
            return json.loads(match.group(0))
        except json.JSONDecodeError:
            return {}
    return {}


def normalize_llm_text(value):
    if isinstance(value, list):
        return '\n\n'.join(str(item).strip() for item in value if str(item).strip())
    if isinstance(value, dict):
        return json.dumps(value, indent=2)
    return str(value or '').strip()


def blank_llm_translation(status='Not requested'):
    return {
        'suggested_power_bi_type': '',
        'suggested_power_bi_expression_action': '',
        'llm_notes': '',
        'needs_manual_review': 'Yes' if status != 'Not requested' else '',
        'review_reason': status,
        'confidence': '',
    }


def translate_expression_with_llm(client, model, expression_type, expression, context):
    if not expression:
        return blank_llm_translation('No raw expression to translate')

    cache_key = (model, expression_type, expression, context)
    cache = st.session_state.setdefault('llm_translation_cache', {})
    if cache_key in cache:
        return cache[cache_key]

    prompt = f"""
You translate Tableau migration logic into Power BI guidance.

Return only valid JSON with these exact keys:
suggested_power_bi_type
suggested_power_bi_expression_action
llm_notes
needs_manual_review
review_reason
confidence

Rules:
- The value for suggested_power_bi_expression_action must be one plain multiline string (DAX or Power Query M).
- Do not return it as a JSON array, Python list, bullet list, or numbered list.
- Prefer a DAX measure for aggregate-level calcs (SUM/AVG/COUNTD-based).
- Prefer a DAX calculated column for row-level, non-aggregate calcs.
- Prefer Power Query M for source shaping, filters before load, merges, and cleanup.
- For LOD expressions (FIXED/INCLUDE/EXCLUDE), be precise: FIXED -> CALCULATE + ALLEXCEPT,
  EXCLUDE -> CALCULATE + ALL(dim), INCLUDE has no context-free DAX equivalent - flag that clearly.
- For table calculations (WINDOW_/RUNNING_/RANK/etc.), explain the DAX window-function equivalent
  and flag that partition/addressing must be confirmed against the original view.
- For relationships/joins, explain the likely Power BI model relationship or Power Query merge,
  and flag that Tableau relationships do not declare cardinality.
- Keep table and column names as placeholders when the exact Power BI model name is unknown.
- Be honest about uncertainty. The result is a migration suggestion, not final production code.
- Set needs_manual_review to "Yes".
- Set confidence to High, Medium, or Low.

Expression type: {expression_type}
Context: {context}
Raw Tableau expression:
{expression}
""".strip()

    try:
        response = client.responses.create(model=model, input=prompt, temperature=0.1)
        data = parse_llm_json(response.output_text)
        result = {
            'suggested_power_bi_type': normalize_llm_text(data.get('suggested_power_bi_type', '')),
            'suggested_power_bi_expression_action': normalize_llm_text(data.get('suggested_power_bi_expression_action', '')),
            'llm_notes': normalize_llm_text(data.get('llm_notes', '')),
            'needs_manual_review': normalize_llm_text(data.get('needs_manual_review', 'Yes')) or 'Yes',
            'review_reason': normalize_llm_text(data.get('review_reason', '')) or 'Review before using in Power BI.',
            'confidence': normalize_llm_text(data.get('confidence', '')),
        }
    except Exception as e:
        result = blank_llm_translation(f'LLM error: {e}')

    cache[cache_key] = result
    return result


def generate_ai_build_plan(client, model, payload):
    """Single on-demand API call, with automatic retry on truncation.
    Returns (BuildPlan | None, error_str | None)."""
    user_prompt = json.dumps(payload, ensure_ascii=False)
    complexity = _COMPLEXITY_ALIASES.get(payload.get('complexity_level', 'Medium'), payload.get('complexity_level', 'Medium'))
    base_budget = {
        'Very Small': 1500, 'Small': 2200, 'Medium': 3500,
        'Large': 5500, 'Very Complex': 7000,
    }.get(complexity, 3500)

    max_budget_cap = 16000
    output_budget = base_budget
    last_raw_text = ''

    while True:
        try:
            response = client.responses.create(
                model=model,
                input=[
                    {"role": "system", "content": AI_BUILD_PLAN_SYSTEM_PROMPT},
                    {"role": "user", "content": user_prompt},
                ],
                temperature=0.1,
                max_output_tokens=output_budget,
            )
        except Exception as e:
            return None, f'LLM error: {e}'

        last_raw_text = response.output_text or ''
        was_truncated = getattr(response, 'status', None) == 'incomplete'
        raw = parse_llm_json(last_raw_text)

        if not raw and was_truncated and output_budget < max_budget_cap:
            output_budget = min(output_budget * 2, max_budget_cap)
            continue

        if not raw:
            reason = (
                ' (response was truncated before completing JSON - '
                'consider shrinking the payload or raising max_output_tokens further)'
                if was_truncated else ''
            )
            return None, f'Model did not return parseable JSON{reason}. Raw response (truncated): {last_raw_text[:500]}'

        raw = normalize_power_bi_layers(raw)

        try:
            plan = BuildPlan.model_validate(raw)
        except ValidationError as ve:
            return None, f'AI response failed schema validation: {ve}'

        allowed_sources = set(payload.get('sources', []))
        for rel in payload.get('relationships', []) + payload.get('relationships_needing_cardinality_check', []):
            for side_key in ('left', 'right'):
                table_part = rel.get(side_key, '').split('[')[0].strip()
                if table_part:
                    allowed_sources.add(table_part)

        allowed_fields = set(payload.get('shown_fields', [])) | {
            item['field'] for item in payload.get('unresolved_formulas', [])
        } | {item['field'] for item in payload.get('lod_translations', [])} | {
            c for c in payload.get('numeric_measure_candidates', [])
        }
        for item in payload.get('unresolved_formulas', []):
            allowed_fields.update(r for r in item.get('references', []))
        for rel in payload.get('relationships', []) + payload.get('relationships_needing_cardinality_check', []):
            for side_key in ('left', 'right'):
                m = re.search(r'\[([^\]]+)\]', rel.get(side_key, ''))
                if m:
                    allowed_fields.add(m.group(1).strip())
        for f in payload.get('filters_sample', []) + payload.get('flagged_filters', []):
            if f.get('field'):
                allowed_fields.add(f['field'])

        plan = strip_unverified_steps(plan, allowed_sources, allowed_fields)
        return plan, None


# ===========================================================================
# EXCEL STYLING HELPERS
# ===========================================================================

def _fill(bg):
    return PatternFill('solid', start_color=bg)


def _font(fg, bold=False):
    return Font(name='Arial', size=9, color=fg, bold=bold)


def _excel_safe_value(val):
    if isinstance(val, str) and len(val) > EXCEL_CELL_LIMIT:
        return val[:EXCEL_CELL_LIMIT - 20] + ' ... [truncated]'
    return val


def _write_cell(ws, row, col, val, bg='FFFFFF', fg='000000', bold=False):
    c = ws.cell(row=row, column=col, value=_excel_safe_value(val))
    if FAST_DATA_FORMATTING and row > 1:
        if bold:
            c.font = _font(fg, bold)
        c.alignment = WRAP_TOP
        return
    c.fill = _fill(bg)
    c.font = _font(fg, bold)
    c.alignment = WRAP_TOP
    c.border = BORDER


def setup_sheet(ws, headers, widths):
    hbg, hfg = COLORS['header']
    for col, (h, w) in enumerate(zip(headers, widths), 1):
        _write_cell(ws, 1, col, h, hbg, hfg, bold=True)
        ws.column_dimensions[get_column_letter(col)].width = w
    ws.row_dimensions[1].height = 22
    ws.freeze_panes = "A2"


def _append_row_values(ws, values, bg='FFFFFF', fg='000000', bold=False):
    cells = []
    for value in values:
        cell = WriteOnlyCell(ws, value=_excel_safe_value(value))
        cell.font = _font(fg, bold)
        cell.alignment = WRAP_TOP
        cell.fill = _fill(bg)
        cell.border = BORDER
        cells.append(cell)
    ws.append(cells)


def _setup_write_only_sheet(wb, title, headers, widths):
    ws = wb.create_sheet(title)
    hbg, hfg = COLORS['header']
    _append_row_values(ws, headers, hbg, hfg, bold=True)
    for col, width in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(col)].width = width
    return ws


# ===========================================================================
# TABLEAU .twb / .twbx PARSER
# ===========================================================================

def load_twb_text(uploaded_file_bytes: bytes, file_name: str) -> str:
    """Returns the raw .twb XML text regardless of whether the upload was
    a .twb or a .twbx."""
    suffix = Path(file_name).suffix.lower()
    if suffix == '.twb':
        return uploaded_file_bytes.decode('utf-8', errors='replace')
    if suffix == '.twbx':
        with zipfile.ZipFile(BytesIO(uploaded_file_bytes), 'r') as zf:
            twb_names = [n for n in zf.namelist() if n.lower().endswith('.twb')]
            if not twb_names:
                raise ValueError(f"No .twb file found inside {file_name}")
            return zf.read(twb_names[0]).decode('utf-8', errors='replace')
    raise ValueError(f"Unsupported file type: {file_name} (expected .twb or .twbx)")


def collect_twbx_from_zip(uploaded_zip_bytes: bytes):
    """For a ZIP-of-workbooks upload: finds every .twb/.twbx inside and
    returns [(folder, file_name, twb_text), ...]."""
    results = []
    with zipfile.ZipFile(BytesIO(uploaded_zip_bytes), 'r') as zf:
        for name in zf.namelist():
            lower = name.lower()
            if lower.endswith('.twb'):
                content = zf.read(name).decode('utf-8', errors='replace')
                results.append((str(Path(name).parent), Path(name).name, content))
            elif lower.endswith('.twbx'):
                inner_bytes = zf.read(name)
                content = load_twb_text(inner_bytes, name)
                results.append((str(Path(name).parent), Path(name).name, content))
    return results


def _strip_brackets(name):
    if name is None:
        return ''
    return name.strip().lstrip('[').rstrip(']')


def _local_field(qualified_name):
    if not qualified_name:
        return ''
    parts = re.findall(r'\[([^\]]+)\]', qualified_name)
    last = parts[-1] if parts else qualified_name
    last = re.sub(r'^(usr|none|ctd|qk|nk|yr|qr):', '', last)
    last = last.split(':')[0]
    return last


def extract_datasources(root):
    """Only reads the workbook-level <datasources> block (direct child of
    <workbook>). Worksheets also carry their own <datasources> reference
    list - that must NOT be swept up here.
    A federated datasource can carry two connections for the same data:
    the live one and a 'hyper' extract cache - prefer the live one."""
    datasources = []
    top_level = root.find('datasources')
    if top_level is None:
        return datasources

    for ds in top_level.findall('datasource'):
        ds_name = ds.get('name', '')
        ds_caption = ds.get('caption', ds_name)

        connection_elements = ds.findall('.//connection')
        live_conns = [c for c in connection_elements if c.get('class') != 'hyper']
        source_conns = live_conns if live_conns else connection_elements

        tables = []
        seen_friendly_names = set()
        for conn in source_conns:
            for rel in conn.findall('.//relation'):
                if rel.get('type') != 'table':
                    continue
                friendly_name = rel.get('name', '')
                table_name = rel.get('table', friendly_name)
                if friendly_name in seen_friendly_names:
                    continue
                seen_friendly_names.add(friendly_name)
                tables.append({
                    'table_name': table_name,
                    'friendly_name': friendly_name,
                    'connection_ref': rel.get('connection', ''),
                    'connection_class': conn.get('class', ''),
                })

        connections = []
        for named_conn in ds.findall('.//named-connections/named-connection'):
            inner = named_conn.find('connection')
            if inner is not None:
                connections.append({
                    'name': named_conn.get('name', ''),
                    'caption': named_conn.get('caption', ''),
                    'class': inner.get('class', ''),
                    'server': inner.get('server', ''),
                    'dbname': inner.get('dbname', inner.get('directory', '')),
                })

        datasources.append({
            'name': ds_name, 'caption': ds_caption,
            'tables': tables, 'connections': connections,
        })

    return datasources


def extract_relationships(root):
    """Handles BOTH join styles: modern noodle-style <relationship> and
    classic SQL-style <relation type='join'>."""
    rows = []

    for rel in root.findall('.//relationships/relationship'):
        expr = rel.find('expression')
        if expr is None:
            continue
        operands = expr.findall('expression')
        if len(operands) != 2:
            continue
        left_field = _local_field(operands[0].get('op', ''))
        right_field = _local_field(operands[1].get('op', ''))
        first_ep = rel.find('first-end-point')
        second_ep = rel.find('second-end-point')
        rows.append({
            'left_object': (first_ep.get('object-id', '') if first_ep is not None else ''),
            'left_field': left_field,
            'right_object': (second_ep.get('object-id', '') if second_ep is not None else ''),
            'right_field': right_field,
            'join_type': expr.get('op', '='),
            'style': 'relationship (logical/noodle)',
            'declares_cardinality': False,
        })

    for rel in root.findall('.//connection//relation'):
        if rel.get('type') != 'join':
            continue
        join_type = rel.get('join', 'inner')
        clause = rel.find('clause')
        children = list(rel)
        left_table = children[0].get('table', children[0].get('name', '')) if len(children) > 0 else ''
        right_table = children[1].get('table', children[1].get('name', '')) if len(children) > 1 else ''
        if clause is not None:
            expr = clause.find('expression')
            if expr is not None:
                operands = expr.findall('expression')
                if len(operands) == 2:
                    rows.append({
                        'left_object': left_table,
                        'left_field': _local_field(operands[0].get('op', '')),
                        'right_object': right_table,
                        'right_field': _local_field(operands[1].get('op', '')),
                        'join_type': join_type,
                        'style': 'classic relation join',
                        'declares_cardinality': False,
                    })

    return rows


def extract_calculated_fields(root):
    """Returns dict keyed by internal field name, deduped, each with:
    caption, formula_raw, formula_resolved, datatype, role, is_table_calc,
    is_lod."""
    calcs = {}
    for col in root.findall('.//column'):
        calc = col.find('calculation')
        if calc is None or calc.get('class') != 'tableau':
            continue
        internal_name = _strip_brackets(col.get('name', ''))
        if not internal_name or internal_name in calcs:
            continue
        formula = calc.get('formula', '') or ''
        calcs[internal_name] = {
            'internal_name': internal_name,
            'caption': col.get('caption', internal_name),
            'datatype': col.get('datatype', ''),
            'role': col.get('role', ''),
            'formula_raw': formula,
            'is_table_calc': bool(re.search(
                r'\b(WINDOW_|RUNNING_|LOOKUP\(|INDEX\(|RANK|TOTAL\(|FIRST\(|LAST\()',
                formula, re.IGNORECASE
            )),
            'is_lod': bool(re.search(r'\{\s*(FIXED|INCLUDE|EXCLUDE)', formula, re.IGNORECASE)),
        }

    name_to_caption = {name: c['caption'] for name, c in calcs.items()}

    def resolve(formula):
        def _sub(m):
            ref = _strip_brackets(m.group(0))
            return f"[{name_to_caption.get(ref, ref)}]"
        return re.sub(r'\[[^\]]+\]', _sub, formula)

    for c in calcs.values():
        c['formula_resolved'] = resolve(c['formula_raw'])
        c['referenced_calcs'] = [
            name_to_caption.get(_strip_brackets(ref), _strip_brackets(ref))
            for ref in re.findall(r'\[([^\]]+)\]', c['formula_raw'])
            if _strip_brackets(ref) in calcs
        ]

    return calcs


def extract_filters(root, calc_name_to_caption):
    """Filters live in <shared-views><shared-view> (report-level) and in
    <worksheets><worksheet>...<filter> (worksheet-only). A [:Measure Names]
    filter is a different kind entirely - it controls which measures are
    swappable, not a row filter - and must be extracted separately."""
    rows = []

    def _measure_names_row(filt, scope, owner):
        member_names = []
        for member in filt.findall('.//groupfilter[@function="member"]'):
            raw_member = member.get('member', '')
            internal = _local_field(raw_member.strip('"'))
            member_names.append(calc_name_to_caption.get(internal, internal))
        return {
            'scope': scope, 'owner': owner, 'filter_class': 'measure-names',
            'field': '(Measure Names filter)', 'function': 'measure-selection',
            'raw_column': filt.get('column', ''),
            'measures_shown': sorted(dict.fromkeys(member_names)),
            'power_bi': (
                'Not a row filter. Controls which measures are swappable/visible in '
                'this view. Rebuild with a Field Parameter (Power BI "Measure swap") '
                'containing exactly these measures: ' + ', '.join(sorted(dict.fromkeys(member_names)))
            ),
        }

    def _row(filt, scope, owner):
        column = filt.get('column', '')
        if column.endswith('[:Measure Names]'):
            return _measure_names_row(filt, scope, owner)
        field = _local_field(column)
        field_display = calc_name_to_caption.get(field, field)
        groupfilter = filt.find('groupfilter')
        function = groupfilter.get('function', '') if groupfilter is not None else ''
        return {
            'scope': scope, 'owner': owner, 'filter_class': filt.get('class', ''),
            'field': field_display, 'function': function, 'raw_column': column,
            'measures_shown': [],
            'power_bi': 'Report-level filter or slicer (applies to every visual)'
                        if scope.startswith('Shared')
                        else 'Visual-level filter on this specific visual',
        }

    for shared_view in root.findall('.//shared-views/shared-view'):
        owner = shared_view.get('name', '')
        for filt in shared_view.findall('filter'):
            rows.append(_row(filt, 'Shared (report-level)', owner))

    for ws in root.findall('.//worksheets/worksheet'):
        owner = ws.get('name', '')
        for filt in ws.findall('.//table/view/filter'):
            rows.append(_row(filt, 'Worksheet-only', owner))

    return rows


def suggest_power_bi_visual(mark_type):
    classes = [c.strip() for c in mark_type.split('+')] if mark_type else []
    if len(classes) > 1:
        return f"Combo chart (Power BI Line and clustered column / combo visual) - combines {mark_type}"
    single = classes[0] if classes else ''
    return {
        'Bar': 'Bar or column chart', 'Line': 'Line chart', 'Area': 'Area chart',
        'Circle': 'Scatter chart', 'Square': 'Scatter chart (square markers)',
        'Shape': 'Scatter chart with custom shapes, or icon-based visual',
        'Text': 'Table, card, or matrix (text-based view)',
        'Pie': 'Pie or donut chart', 'Map': 'Map visual (Azure Maps / Shape map)',
        'GanttBar': 'Gantt-style visual (custom or AppSource Gantt)',
        'Polygon': 'Filled/shape map or custom polygon visual',
        'Density': 'Heat map / density visual',
        'Automatic': 'Review worksheet - Tableau auto-selected the mark; check the actual rendered chart type',
        '': 'Could not determine mark type - review worksheet manually',
    }.get(single, f'Unrecognized mark type "{single}" - review manually')


def extract_worksheets(root, calc_name_to_caption):
    rows = []
    for ws in root.findall('.//worksheets/worksheet'):
        name = ws.get('name', '')
        fields_used = []
        for col in ws.findall('.//datasource-dependencies/column'):
            internal = _strip_brackets(col.get('name', ''))
            caption = col.get('caption') or calc_name_to_caption.get(internal, internal)
            fields_used.append(caption)

        mark_classes = [m.get('class', '') for m in ws.findall('.//table/panes/pane/mark')]
        mark_type = ' + '.join(dict.fromkeys(c for c in mark_classes if c))

        rows.append({
            'worksheet_name': name,
            'fields_used': sorted(dict.fromkeys(fields_used)),
            'mark_type': mark_type,
            'power_bi_visual': suggest_power_bi_visual(mark_type),
        })
    return rows


def extract_dashboards(root, worksheet_names):
    rows = []
    for dash in root.findall('.//dashboards/dashboard'):
        dash_name = dash.get('name', '')
        placed = []
        for zone in dash.findall('.//zone'):
            z_name = zone.get('name', '')
            z_type = zone.get('type-v2', '')
            if z_name in worksheet_names and z_type == '':
                placed.append(z_name)
        rows.append({'dashboard_name': dash_name, 'worksheets_placed': sorted(dict.fromkeys(placed))})
    return rows


ACTION_KIND_MAP = {
    'filter-action': 'Filter action', 'highlight-action': 'Highlight action',
    'edit-parameter-action': 'Parameter action', 'url-action': 'URL / navigation action',
    'go-to-sheet-action': 'Go to sheet (drillthrough-like)',
}


def suggest_action_power_bi_equivalent(kind):
    return {
        'Filter action': 'Cross-report/visual-level filter interaction (default Power BI cross-filtering, or a Bookmark+Filter for custom behavior).',
        'Highlight action': 'Visual-level highlighting (Power BI does this by default on click; custom highlight sets may need a Bookmark.',
        'Parameter action': 'Power BI "What-if" parameter updated via a Field parameter + Bookmark, or dynamic measure selection.',
        'URL / navigation action': 'Power BI Q&A/Web URL button or "Web browser" navigation action on click.',
        'Go to sheet (drillthrough-like)': 'Power BI Drillthrough page.',
    }.get(kind, 'Review manually - no direct Power BI equivalent identified.')


def extract_actions(root):
    rows = []
    actions_el = root.find('.//actions')
    if actions_el is None:
        return rows
    for action in list(actions_el):
        tag = action.tag
        kind = ACTION_KIND_MAP.get(tag, tag)
        source = action.find('source')
        rows.append({
            'action_name': action.get('caption', action.get('name', '')),
            'kind': kind,
            'source_worksheet': source.get('worksheet', '') if source is not None else '',
            'activation': (action.find('activation').get('type', '')
                           if action.find('activation') is not None else ''),
            'power_bi_equivalent': suggest_action_power_bi_equivalent(kind),
        })
    return rows


# --- LOD (Level of Detail) -> DAX translation -----------------------------

_TABLEAU_TO_DAX_AGG = {
    'SUM': 'SUM', 'AVG': 'AVERAGE', 'MIN': 'MIN', 'MAX': 'MAX',
    'COUNT': 'COUNT', 'COUNTD': 'DISTINCTCOUNT', 'MEDIAN': 'MEDIAN',
}

LOD_PATTERN = re.compile(
    r'\{\s*(FIXED|INCLUDE|EXCLUDE)\s*(.*?)\s*:\s*(.*?)\s*\}',
    re.IGNORECASE | re.DOTALL,
)


def _tableau_expr_to_dax(expr, table_name):
    result = re.sub(r'\[([^\]]+)\]', lambda m: f'{table_name}[{m.group(1)}]', expr)

    def _agg_sub(m):
        return f'{_TABLEAU_TO_DAX_AGG.get(m.group(1).upper(), m.group(1).upper())}('

    result = re.sub(r'\b(SUM|AVG|MIN|MAX|COUNT|COUNTD|MEDIAN)\s*\(', _agg_sub, result, flags=re.IGNORECASE)
    return result


def parse_lod_expressions(formula):
    found = []
    for match in LOD_PATTERN.finditer(formula):
        lod_type = match.group(1).upper()
        dims_raw = match.group(2).strip()
        agg_expr = match.group(3).strip()
        dims = [d.strip().strip('[]') for d in dims_raw.split(',') if d.strip()]
        found.append({'lod_type': lod_type, 'dims': dims, 'agg_expr_raw': agg_expr})
    return found


def generate_lod_dax(lod_type, dims, agg_expr_raw, table_name='Table'):
    dax_agg = _tableau_expr_to_dax(agg_expr_raw, table_name)
    dims_dax = [f'{table_name}[{d}]' for d in dims]

    if lod_type == 'FIXED':
        dax = f"CALCULATE(\n    {dax_agg},\n    ALLEXCEPT({table_name}, {', '.join(dims_dax)})\n)"
        caveat = (
            "FIXED ignores the view's filters/grouping except explicit context filters. "
            "If the original Tableau view also used context filters, those must be applied "
            "BEFORE this CALCULATE or the fixed total will include rows Tableau would have excluded."
        )
        return dax, caveat

    if lod_type == 'EXCLUDE':
        all_clauses = ', '.join(f'ALL({d})' for d in dims_dax)
        return f"CALCULATE(\n    {dax_agg},\n    {all_clauses}\n)", None

    if lod_type == 'INCLUDE':
        dims_str = ', '.join(dims_dax)
        dax = f"SUMX(\n    SUMMARIZE({table_name}, {dims_str}),\n    {dax_agg}\n)"
        caveat = (
            "INCLUDE's result depends on which dimensions are present in the Tableau view "
            "(it adds extra granularity before the view re-aggregates). Confirm how the ORIGINAL "
            "VIEW re-aggregates this value and wrap this expression accordingly before use."
        )
        return dax, caveat

    return f"-- Unrecognized LOD type: {lod_type}", "Could not classify this LOD expression - review manually."


def _run_lod_self_tests():
    fixed = parse_lod_expressions('{ FIXED [Region],[Category] : SUM([Sales]) }')
    assert fixed[0]['lod_type'] == 'FIXED'
    dax, _ = generate_lod_dax('FIXED', ['Region', 'Category'], 'SUM([Sales])', 'Orders')
    assert 'ALLEXCEPT(Orders, Orders[Region], Orders[Category])' in dax
    dax2, _ = generate_lod_dax('EXCLUDE', ['Customer'], 'AVG([Order Value])', 'Orders')
    assert 'ALL(Orders[Customer])' in dax2 and 'AVERAGE(Orders[Order Value])' in dax2
    dax3, caveat3 = generate_lod_dax('INCLUDE', ['Order ID'], 'COUNTD([Product])', 'Orders')
    assert 'DISTINCTCOUNT(Orders[Product])' in dax3 and caveat3 is not None


_run_lod_self_tests()


def classify_tableau_formula(calc):
    formula = calc['formula_raw'].upper()
    types = []
    if calc['is_lod']:
        types.append('Level of Detail (FIXED/INCLUDE/EXCLUDE)')
    if calc['is_table_calc']:
        types.append('Table calculation (WINDOW_/RUNNING_/RANK/LOOKUP/INDEX)')
    if re.search(r'\bIF\b.+\bTHEN\b', formula):
        types.append('Conditional IF')
    if re.search(r'\bCASE\b.+\bWHEN\b', formula):
        types.append('CASE/Mapping')
    if re.search(r'\b(SUM|AVG|COUNT|COUNTD|MIN|MAX|MEDIAN)\s*\(', formula):
        types.append('Aggregation')
    if re.search(r'\b(DATEADD|DATEDIFF|DATEPART|DATETRUNC|TODAY|NOW)\s*\(', formula):
        types.append('Date/Time')
    if re.search(r'\b(LEFT|RIGHT|MID|CONTAINS|REPLACE|SPLIT|TRIM|UPPER|LOWER)\s*\(', formula):
        types.append('String')
    if re.search(r'(?<![<>=])[-+*/](?![<>=])', formula) and not types:
        types.append('Arithmetic')
    return ' + '.join(types) if types else 'Direct/Other'


def suggest_tableau_formula_rebuild(formula_type):
    if 'Level of Detail' in formula_type:
        return 'Rebuild as a DAX measure using CALCULATE with explicit filter context (FIXED ~ ALLEXCEPT, INCLUDE/EXCLUDE ~ ALL/removing groupby columns).'
    if 'Table calculation' in formula_type:
        return 'Rebuild as a DAX measure using window functions (RANKX, running-total pattern, or WINDOW-family DAX functions). Confirm partition/addressing against the original view.'
    if 'CASE/Mapping' in formula_type:
        return 'Use DAX SWITCH() or a Power Query conditional/lookup merge.'
    if 'Conditional IF' in formula_type:
        return 'Use DAX IF/SWITCH or a Power Query conditional column.'
    if 'Aggregation' in formula_type:
        return 'Use a DAX measure (SUM/AVERAGE/COUNT/DISTINCTCOUNT/MIN/MAX/MEDIAN).'
    if 'Date/Time' in formula_type:
        return 'Use DAX date functions or Power Query date transformations.'
    if 'String' in formula_type:
        return 'Use DAX text functions or Power Query text transformations.'
    if 'Arithmetic' in formula_type:
        return 'Use a DAX calculated column/measure or a Power Query custom column.'
    return 'Review expression and rebuild as DAX or Power Query M.'


def parse_twb(twb_text: str) -> dict:
    root = ET.fromstring(twb_text)

    datasources = extract_datasources(root)
    relationships = extract_relationships(root)
    calcs = extract_calculated_fields(root)
    calc_name_to_caption = {name: c['caption'] for name, c in calcs.items()}
    filters = extract_filters(root, calc_name_to_caption)
    worksheets = extract_worksheets(root, calc_name_to_caption)
    worksheet_names = {w['worksheet_name'] for w in worksheets}
    dashboards = extract_dashboards(root, worksheet_names)
    actions = extract_actions(root)

    formula_mapping = []
    lod_translations = []
    for internal_name, calc in calcs.items():
        formula_type = classify_tableau_formula(calc)

        if calc['is_lod']:
            for lod in parse_lod_expressions(calc['formula_raw']):
                dax, caveat = generate_lod_dax(
                    lod['lod_type'], lod['dims'], lod['agg_expr_raw'], table_name='FactTable',
                )
                lod_translations.append({
                    'field': calc['caption'], 'lod_type': lod['lod_type'],
                    'dims': lod['dims'], 'agg_expr_raw': lod['agg_expr_raw'],
                    'dax': dax, 'caveat': caveat,
                })

        formula_mapping.append({
            'field': calc['caption'], 'internal_name': internal_name,
            'datatype': calc['datatype'], 'role': calc['role'],
            'formula_type': formula_type, 'formula_raw': calc['formula_raw'],
            'formula_resolved': calc['formula_resolved'],
            'referenced_calcs': calc['referenced_calcs'],
            'power_bi': suggest_tableau_formula_rebuild(formula_type),
        })

    return {
        'datasources': datasources, 'relationships': relationships,
        'calculated_fields': calcs, 'formula_mapping': formula_mapping,
        'lod_translations': lod_translations, 'filters': filters,
        'worksheets': worksheets, 'dashboards': dashboards, 'actions': actions,
    }


# ===========================================================================
# DUPLICATE WORKBOOK DETECTION
# ===========================================================================

def _norm_set(values):
    return frozenset(str(v).strip().upper() for v in values if str(v or '').strip())


def compute_workbook_profile(parsed):
    tables = _norm_set(t['friendly_name'] for ds in parsed['datasources'] for t in ds['tables'])
    fields = _norm_set(c['caption'] for c in parsed['calculated_fields'].values())
    formulas = frozenset(
        (row['field'].upper(), re.sub(r'\s+', ' ', row['formula_raw']).upper())
        for row in parsed['formula_mapping']
    )
    filters = frozenset(
        (f['field'].upper(), f['function'].upper())
        for f in parsed['filters'] if f['filter_class'] != 'measure-names'
    )
    joins = frozenset(
        (r['left_field'].upper(), r['right_field'].upper(), r['join_type'].upper())
        for r in parsed['relationships']
    )
    worksheet_names = _norm_set(w['worksheet_name'] for w in parsed['worksheets'])

    if not tables and not fields and not worksheet_names:
        return None

    return {
        'tables': tables, 'fields': fields, 'formulas': formulas,
        'filters': filters, 'joins': joins, 'worksheet_names': worksheet_names,
        'all_fields': tables | fields,
    }


def _profile_exact_key(profile):
    if profile is None:
        return None
    return (profile['tables'], profile['fields'], profile['formulas'], profile['filters'], profile['joins'])


def _field_similarity(a, b):
    if not a and not b:
        return 1.0
    if not a or not b:
        return 0.0
    return len(a & b) / len(a | b)


def build_workbook_duplicate_analysis(workbook_results):
    records = []
    exact_groups = defaultdict(list)
    table_groups = defaultdict(list)

    for folder, file_name, parsed in workbook_results:
        profile = compute_workbook_profile(parsed) if parsed is not None else None
        record = {
            'folder': folder, 'file_name': file_name, 'parsed': parsed, 'profile': profile,
            'category': 'Unparsed' if profile is None else 'Unique',
            'group_id': '', 'group_size': 1,
            'exact_matches': [], 'near_matches': [], 'same_source_matches': [],
            'best_similarity': 0.0, 'difference_summary': '',
        }
        records.append(record)
        if profile is not None:
            exact_groups[_profile_exact_key(profile)].append(record)
            table_groups[profile['tables']].append(record)

    group_counter = 1
    for members in exact_groups.values():
        if len(members) <= 1:
            continue
        group_id = f'Exact {group_counter}'
        group_counter += 1
        names = [m['file_name'] for m in members]
        for record in members:
            record['category'] = 'Exact Duplicate'
            record['group_id'] = group_id
            record['group_size'] = len(members)
            record['exact_matches'] = [n for n in names if n != record['file_name']]
            record['best_similarity'] = 1.0
            record['difference_summary'] = 'Exact same tables, fields, formulas, filters, and joins'

    for table_key, members in table_groups.items():
        if not table_key or len(members) <= 1:
            continue
        for record in members:
            if record['category'] == 'Exact Duplicate':
                continue
            best_near, same_source, best_sim, diffs = [], [], 0.0, []
            for other in members:
                if other is record:
                    continue
                sim = _field_similarity(record['profile']['all_fields'], other['profile']['all_fields'])
                best_sim = max(best_sim, sim)
                if sim >= 0.80:
                    if (record['profile']['filters'] != other['profile']['filters']
                            or record['profile']['formulas'] != other['profile']['formulas']
                            or record['profile']['joins'] != other['profile']['joins']):
                        best_near.append(other['file_name'])
                        for label, key in [('fields', 'all_fields'), ('filters', 'filters'),
                                            ('formulas', 'formulas'), ('joins', 'joins')]:
                            if record['profile'][key] != other['profile'][key]:
                                diffs.append(label)
                    else:
                        same_source.append(other['file_name'])
                else:
                    same_source.append(other['file_name'])

            if best_near:
                record['category'] = 'Near Duplicate'
                record['near_matches'] = sorted(set(best_near))
                record['group_id'] = f'Near Source {abs(hash(table_key)) % 100000}'
                record['group_size'] = len(best_near) + 1
                record['difference_summary'] = '; '.join(sorted(set(diffs))) or '80%+ same fields with logic differences'
            elif same_source:
                record['category'] = 'Same Data Source'
                record['same_source_matches'] = sorted(set(same_source))
                record['group_id'] = f'Source {abs(hash(table_key)) % 100000}'
                record['group_size'] = len(same_source) + 1
                record['difference_summary'] = 'Same source tables with different fields or logic'
            record['best_similarity'] = max(record['best_similarity'], best_sim)

    for record in records:
        if record['category'] == 'Unique':
            record['difference_summary'] = 'No meaningful match found'
        elif record['category'] == 'Unparsed':
            record['difference_summary'] = 'Could not parse enough tables or fields'

    return records


def duplicate_counts(duplicate_records):
    counts = Counter(record['category'] for record in duplicate_records)
    return {
        'exact': counts.get('Exact Duplicate', 0), 'near': counts.get('Near Duplicate', 0),
        'same_source': counts.get('Same Data Source', 0), 'unique': counts.get('Unique', 0),
        'unparsed': counts.get('Unparsed', 0),
    }


# ===========================================================================
# CONTENT USAGE FILE (Tableau Server/Cloud usage export) — optional,
# equivalent role to the WebFOCUS Resource Analyzer file: lets the user
# filter/match uploaded workbooks against an actual usage/active-content list.
# ===========================================================================

def normalize_workbook_name(value):
    if value is None:
        return ''
    text = str(value).strip()
    if not text or text.lower() == 'nan':
        return ''
    text = text.replace('\\', '/').split('/')[-1].strip()
    for ext in ('.twbx', '.twb'):
        if text.lower().endswith(ext):
            text = text[: -len(ext)]
    text = re.sub(r'[^A-Za-z0-9_.$# -]+', '', text)
    return text.upper().strip()


def extract_workbook_tokens_from_text(value):
    if value is None:
        return set()
    text = str(value).strip()
    if not text or text.lower() == 'nan':
        return set()
    tokens = set()
    twbx_matches = re.findall(r'([A-Za-z0-9_.$#/ -]+\.twbx?)', text, flags=re.IGNORECASE)
    for item in twbx_matches:
        normalized = normalize_workbook_name(item)
        if normalized:
            tokens.add(normalized)
    if not tokens:
        possible = re.findall(r'\b[A-Za-z0-9_.$# -]{3,}\b', text)
        for word in possible:
            cleaned = normalize_workbook_name(word)
            if cleaned:
                tokens.add(cleaned)
    return tokens


def read_usage_file(uploaded_file):
    """Reads a Content Usage export (xlsx/csv) and extracts workbook-name
    tokens, mirroring read_resource_analyzer_file() from the WebFOCUS app."""
    file_name = uploaded_file.name.lower()
    workbook_names = set()
    uploaded_file.seek(0)
    try:
        if file_name.endswith('.csv'):
            df_map = {'Usage': pd.read_csv(uploaded_file, dtype=str, header=None)}
        else:
            df_map = pd.read_excel(uploaded_file, sheet_name=None, dtype=str, header=None)
    except Exception as e:
        st.warning(f"Content usage file could not be read properly: {e}")
        return workbook_names

    for _, df in df_map.items():
        if df is None or df.empty:
            continue
        df = df.fillna('')
        for row_index in range(df.shape[0]):
            for value in df.iloc[row_index].astype(str).tolist():
                value = str(value).strip()
                if not value or value.lower() == 'nan':
                    continue
                for item in extract_workbook_tokens_from_text(value):
                    if item:
                        workbook_names.add(item)

    return workbook_names


def filter_workbooks_by_usage(workbook_items, allowed_workbook_names):
    """workbook_items: [(folder, file_name, text), ...]."""
    matched_items, matched_pairs = [], []
    allowed_clean = {normalize_workbook_name(n) for n in allowed_workbook_names}
    for folder, file_name, content in workbook_items:
        normalized = normalize_workbook_name(file_name)
        if normalized in allowed_clean:
            matched_items.append((folder, file_name, content))
            matched_pairs.append((normalized, file_name))
    return matched_items, matched_pairs


# ===========================================================================
# TDS / TDSX METADATA — published Tableau Data Source files, used to map
# workbook table references to their real underlying database table/
# connection (Tableau's analog to WebFOCUS's .mas/.acx metadata files).
# ===========================================================================

def collect_tds_metadata_from_zip(uploaded_zip_bytes):
    """Reads .tds files from an uploaded ZIP (and unpacks .tdsx the same way
    a .twbx is unpacked). Returns {normalized_table_name: (file_name, table_info)}."""
    tds_map = {}
    with zipfile.ZipFile(BytesIO(uploaded_zip_bytes), 'r') as zf:
        for name in zf.namelist():
            lower = name.lower()
            text = None
            if lower.endswith('.tds'):
                text = zf.read(name).decode('utf-8', errors='replace')
            elif lower.endswith('.tdsx'):
                inner_bytes = zf.read(name)
                with zipfile.ZipFile(BytesIO(inner_bytes), 'r') as inner_zf:
                    tds_names = [n for n in inner_zf.namelist() if n.lower().endswith('.tds')]
                    if tds_names:
                        text = inner_zf.read(tds_names[0]).decode('utf-8', errors='replace')
            if text is None:
                continue
            try:
                root = ET.fromstring(text)
            except ET.ParseError:
                continue
            datasources = extract_datasources(root)
            for ds in datasources:
                for t in ds['tables']:
                    key = normalize_workbook_name(t['friendly_name'])
                    if key and key not in tds_map:
                        tds_map[key] = (Path(name).name, t)
    return tds_map


def build_table_db_mapping(unique_tables, tds_map):
    rows = []
    mapped_count = 0
    needs_review_count = 0
    for table_name in sorted(unique_tables):
        key = normalize_workbook_name(table_name)
        entry = tds_map.get(key)
        if entry:
            file_name, t_info = entry
            rows.append({
                'Unique Table from KashMap': table_name,
                'Actual Database Table': t_info.get('table_name', ''),
                'Connection': t_info.get('connection_class', ''),
                'TDS/TDSX File': file_name,
                'Evidence / Notes': f"Matched via published data source relation for '{t_info.get('friendly_name', '')}'.",
            })
            mapped_count += 1
        else:
            rows.append({
                'Unique Table from KashMap': table_name,
                'Actual Database Table': '', 'Connection': '', 'TDS/TDSX File': '',
                'Evidence / Notes': 'No matching TDS/TDSX metadata file found. Needs review.',
            })
            needs_review_count += 1
    return rows, mapped_count, needs_review_count


# ===========================================================================
# AI BUILD PLAN - payload construction (compact, twb-only)
# ===========================================================================

def _numeric_measure_candidates(worksheets):
    candidates = []
    value_words = r'AMOUNT|COST|SALES|PRICE|REVENUE|SPEND|BOOKING|COUNT|TOTAL|VALUE|RATE'
    helper_words = r'DATE|DAY|OFFSET|RANK|SEQ|SEQUENCE|KEY|CODE|STATUS|FLAG|YEAR|MONTH|WEEK|ID\b'
    for w in worksheets:
        for field in w['fields_used']:
            if re.search(value_words, field, re.IGNORECASE) and not re.search(helper_words, field, re.IGNORECASE):
                candidates.append(field)
    return list(dict.fromkeys(candidates))[:15]


def build_ai_build_plan_payload(workbook_row, parsed):
    """Compact .twb-only payload for the single AI call. Rule-resolved items
    (LOD translations) are summarized as authoritative, so the AI never
    re-derives what the rule engine already computed with certainty."""

    sources = []
    for ds in parsed['datasources']:
        sources.extend(t['friendly_name'] for t in ds['tables'])
    sources = list(dict.fromkeys(sources))[:20]

    relationships = [
        {'left': r['left_field'], 'right': r['right_field'], 'op': r['join_type'], 'style': r['style']}
        for r in parsed['relationships']
    ]
    relationships_needing_cardinality_check = list(relationships)

    shown_fields = []
    for w in parsed['worksheets']:
        shown_fields.extend(w['fields_used'])
    shown_fields = list(dict.fromkeys(shown_fields))[:60]

    lod_translations = [
        {'field': lod['field'], 'lod_type': lod['lod_type'], 'dims': lod['dims'],
         'dax': lod['dax'], 'has_caveat': bool(lod['caveat'])}
        for lod in parsed['lod_translations']
    ]
    lod_fields = {lod['field'] for lod in parsed['lod_translations']}

    unresolved_formulas = []
    for row in parsed['formula_mapping']:
        if row['field'] in lod_fields:
            continue
        risky = any(t in row['formula_type'] for t in [
            'Table calculation', 'Conditional IF', 'CASE/Mapping', 'Date/Time', 'String', 'Arithmetic',
        ])
        if risky:
            unresolved_formulas.append({
                'field': row['field'], 'formula_type': row['formula_type'],
                'raw_formula': row['formula_raw'][:300], 'references': row['referenced_calcs'],
            })
    unresolved_formulas = unresolved_formulas[:15]

    filters_summary, flagged_filters = [], []
    for f in parsed['filters']:
        if f['filter_class'] == 'measure-names':
            flagged_filters.append({'field': '(Measure Names)', 'note': 'Field Parameter, not a slicer', 'measures': f['measures_shown']})
            continue
        filters_summary.append({'scope': f['scope'], 'field': f['field'], 'function': f['function']})
    filters_summary = filters_summary[:15]

    dashboards_sample = [{'name': d['dashboard_name'], 'worksheets': d['worksheets_placed']} for d in parsed['dashboards']]
    actions_sample = [{'kind': a['kind'], 'source': a['source_worksheet']} for a in parsed['actions']]

    return {
        'report_name': str(workbook_row.get('Workbook Name', '')),
        'complexity_level': str(workbook_row.get('Complexity', 'Medium')),
        'sources': sources,
        'shown_fields': shown_fields,
        'relationships': relationships,
        'relationships_needing_cardinality_check': relationships_needing_cardinality_check,
        'filters_sample': filters_summary,
        'flagged_filters': flagged_filters,
        'unresolved_formulas': unresolved_formulas,
        'lod_translations': lod_translations,
        'numeric_measure_candidates': _numeric_measure_candidates(parsed['worksheets']),
        'dashboards_sample': dashboards_sample,
        'actions_sample': actions_sample,
        'constraints': {
            'database_metadata_available': False,
            'may_not_state_primary_keys': True,
            'may_not_state_cardinality': True,
            'tableau_relationships_do_not_declare_cardinality': True,
            'may_not_invent_tables_or_fields': True,
        },
    }


# ===========================================================================
# RULE-BASED POWER QUERY M DRAFT (Tableau's analog to the WebFOCUS SQL View
# draft - Tableau has no HOLD-chain concept, so instead of SQL CTEs this
# generates a Power Query M merge draft from the relationships).
# ===========================================================================

def m_identifier(value, fallback='Table'):
    text = str(value or '').strip()
    text = re.sub(r'[^A-Za-z0-9_ ]+', '_', text)
    text = text.strip('_') or fallback
    return text


def generate_power_query_m_draft(workbook_row, parsed):
    tables = []
    for ds in parsed['datasources']:
        tables.extend(t['friendly_name'] for t in ds['tables'])
    tables = list(dict.fromkeys(tables))

    lines = [
        "// Rule-based Power Query M draft / merge blueprint",
        "// Generated from parsed Tableau workbook structure.",
        "// This is a starting point, not production-ready M code.",
        "// REVIEW: confirm real source connection strings/paths before running.",
        "// REVIEW: Tableau relationships do not declare cardinality - verify",
        "//         join kind (Inner/LeftOuter/etc.) and key uniqueness below.",
        "",
    ]

    if not tables:
        lines.append("// No source tables were detected for this workbook.")
        return '\n'.join(lines)

    for t in tables:
        var = m_identifier(t)
        lines.append(f'{var} = Source{{[Name="{t}"]}}[Data],  // TODO: replace with real source step')

    lines.append('')
    if parsed['relationships']:
        base_table = m_identifier(tables[0])
        prev_var = base_table
        for i, rel in enumerate(parsed['relationships'], start=1):
            merged_var = f"Merged{i}"
            lines.append(
                f'{merged_var} = Table.NestedJoin({prev_var}, {{"{rel["left_field"]}"}}, '
                f'<OtherTableVar>, {{"{rel["right_field"]}"}}, "Joined{i}", JoinKind.LeftOuter),'
                f'  // TODO: confirm join kind - Tableau relationship does not declare cardinality'
            )
            prev_var = merged_var
        lines.append('')

    lines.append("// Final step: expand joined columns, set data types, rename for the semantic model.")
    return '\n'.join(lines)


# ===========================================================================
# VALIDATION CHECKLIST
# ===========================================================================

def build_validation_rows_raw(workbook_name, has_relationships=False, has_lod=False,
                               has_table_calc=False, has_actions=False, has_measure_names_filter=False):
    rows = []

    def add_check(category, check, rule):
        rows.append({'Workbook Name': workbook_name, 'Check Category': category,
                      'Validation Check': check, 'Trigger Rule': rule})

    add_check('Data Validation', 'Compare final row count between Tableau output and Power BI dataset.', 'Always')
    if has_relationships:
        add_check('Data Validation', 'Confirm relationship cardinality and key uniqueness (Tableau relationships do not declare this).', 'Relationship detected')

    add_check('Business Total Validation', 'Validate totals for the main measure/value columns.', 'Always')
    if has_lod:
        add_check('Business Total Validation', 'Validate LOD-derived totals against Tableau at the same grouping level.', 'LOD expression detected')
    if has_table_calc:
        add_check('Business Total Validation', 'Validate table-calculation-derived values (running totals, ranks, window calcs) match Tableau exactly.', 'Table calculation detected')

    add_check('Report Logic Validation', 'Check filters and parameters match the Tableau workbook.', 'Always')
    add_check('Report Logic Validation', 'Confirm final visual type and grouping fields.', 'Always')
    if has_measure_names_filter:
        add_check('Report Logic Validation', 'Confirm the Field Parameter measure-swap list matches the original Measure Names filter exactly.', 'Measure Names filter detected')

    if has_actions:
        add_check('Interaction Validation', 'Confirm drillthrough/cross-filter/parameter actions behave the same as the original Tableau actions.', 'Action detected')

    return rows


# ===========================================================================
# DISPLAY TABLES (flatten parsed dicts into DataFrames)
# ===========================================================================

def workbook_complexity_score(parsed):
    return (
        len(parsed['formula_mapping'])
        + len(parsed['filters'])
        + len(parsed['relationships']) * 2
        + len(parsed['worksheets'])
        + len(parsed['lod_translations']) * 2
    )


def workbook_complexity_label(score):
    if score >= 30:
        return 'High'
    if score >= 14:
        return 'Medium'
    return 'Low'


def build_tables_df(folder, file_name, parsed):
    rows = []
    for ds in parsed['datasources']:
        for t in ds['tables']:
            rows.append({'File Path': folder, 'Workbook Name': file_name, 'Datasource': ds['caption'],
                         'Table': t['friendly_name'], 'Connection Class': t.get('connection_class', '')})
    return pd.DataFrame(rows)


def build_relationships_df(folder, file_name, parsed):
    rows = []
    for r in parsed['relationships']:
        rows.append({'File Path': folder, 'Workbook Name': file_name, 'Left Field': r['left_field'],
                     'Operator': r['join_type'], 'Right Field': r['right_field'], 'Style': r['style']})
    return pd.DataFrame(rows)


def build_formulas_df(folder, file_name, parsed):
    rows = []
    for f in parsed['formula_mapping']:
        rows.append({'File Path': folder, 'Workbook Name': file_name, 'Field': f['field'],
                     'Formula Type': f['formula_type'], 'Raw Formula': f['formula_raw'],
                     'Resolved Formula': f['formula_resolved'], 'References': ', '.join(f['referenced_calcs']),
                     'Power BI Guidance': f['power_bi']})
    return pd.DataFrame(rows)


def build_lod_df(folder, file_name, parsed):
    rows = []
    for lod in parsed['lod_translations']:
        rows.append({'File Path': folder, 'Workbook Name': file_name, 'Field': lod['field'],
                     'LOD Type': lod['lod_type'], 'Dimensions': ', '.join(lod['dims']),
                     'Tableau Expression': lod['agg_expr_raw'], 'Suggested DAX': lod['dax'],
                     'Caveat': lod['caveat'] or ''})
    return pd.DataFrame(rows)


def build_filters_df(folder, file_name, parsed):
    rows = []
    for f in parsed['filters']:
        rows.append({'File Path': folder, 'Workbook Name': file_name, 'Scope': f['scope'],
                     'Owner': f['owner'], 'Field': f['field'], 'Function': f['function'],
                     'Power BI Guidance': f['power_bi']})
    return pd.DataFrame(rows)


def build_worksheets_df(folder, file_name, parsed):
    rows = []
    for w in parsed['worksheets']:
        rows.append({'File Path': folder, 'Workbook Name': file_name, 'Worksheet': w['worksheet_name'],
                     'Mark Type': w['mark_type'], 'Suggested Power BI Visual': w['power_bi_visual'],
                     'Fields Used': ', '.join(w['fields_used'])})
    return pd.DataFrame(rows)


def build_dashboards_df(folder, file_name, parsed):
    rows = []
    for d in parsed['dashboards']:
        rows.append({'File Path': folder, 'Workbook Name': file_name, 'Dashboard': d['dashboard_name'],
                     'Worksheets Placed': ', '.join(d['worksheets_placed'])})
    return pd.DataFrame(rows)


def build_actions_df(folder, file_name, parsed):
    rows = []
    for a in parsed['actions']:
        rows.append({'File Path': folder, 'Workbook Name': file_name, 'Action': a['action_name'],
                     'Kind': a['kind'], 'Source Worksheet': a['source_worksheet'],
                     'Activation': a['activation'], 'Power BI Equivalent': a['power_bi_equivalent']})
    return pd.DataFrame(rows)


def build_duplicates_df(dup_records):
    rows = []
    for r in dup_records:
        rows.append({
            'Duplicate Type': r['category'], 'Group ID': r['group_id'],
            'File Path': r['folder'], 'Workbook Name': r['file_name'],
            'Similarity': f"{r['best_similarity']:.0%}" if r['best_similarity'] else '',
            'Matched With': ', '.join(r['exact_matches'] + r['near_matches'] + r['same_source_matches']),
            'Difference Summary': r['difference_summary'],
        })
    return pd.DataFrame(rows)


def build_display_tables(workbook_results, dup_records, allowed_usage_names, matched_pairs, tds_map):
    dup_lookup = {(r['folder'], r['file_name']): r for r in dup_records}

    overview_rows, validation_rows = [], []
    all_tables, all_rels, all_formulas, all_lods = [], [], [], []
    all_filters, all_worksheets, all_dashboards, all_actions = [], [], [], []

    unique_tables = set()

    for folder, file_name, parsed in workbook_results:
        dup = dup_lookup.get((folder, file_name), {})
        if parsed is None:
            overview_rows.append({
                'File Path': folder, 'Workbook Name': file_name, 'Datasources': 0, 'Tables': 0,
                'Relationships': 0, 'Calculated Fields': 0, 'Filters': 0, 'LOD Translations': 0,
                'Worksheets': 0, 'Dashboards': 0, 'Actions': 0, 'Complexity': 'Manual Review',
                'Duplicate Type': dup.get('category', 'Unparsed'), 'Parse Status': 'Unparsed',
            })
            validation_rows.extend(build_validation_rows_raw(file_name))
            continue

        total_tables = sum(len(ds['tables']) for ds in parsed['datasources'])
        for ds in parsed['datasources']:
            for t in ds['tables']:
                unique_tables.add(t['friendly_name'])

        score = workbook_complexity_score(parsed)
        complexity = workbook_complexity_label(score)
        has_lod = bool(parsed['lod_translations'])
        has_table_calc = any('Table calculation' in f['formula_type'] for f in parsed['formula_mapping'])
        has_measure_names = any(f['filter_class'] == 'measure-names' for f in parsed['filters'])

        overview_rows.append({
            'File Path': folder, 'Workbook Name': file_name,
            'Datasources': len(parsed['datasources']), 'Tables': total_tables,
            'Relationships': len(parsed['relationships']), 'Calculated Fields': len(parsed['calculated_fields']),
            'Filters': len(parsed['filters']), 'LOD Translations': len(parsed['lod_translations']),
            'Worksheets': len(parsed['worksheets']), 'Dashboards': len(parsed['dashboards']),
            'Actions': len(parsed['actions']), 'Complexity': complexity,
            'Duplicate Type': dup.get('category', 'Unique'), 'Parse Status': 'Parsed',
        })

        validation_rows.extend(build_validation_rows_raw(
            file_name, has_relationships=bool(parsed['relationships']), has_lod=has_lod,
            has_table_calc=has_table_calc, has_actions=bool(parsed['actions']),
            has_measure_names_filter=has_measure_names,
        ))

        all_tables.append(build_tables_df(folder, file_name, parsed))
        all_rels.append(build_relationships_df(folder, file_name, parsed))
        all_formulas.append(build_formulas_df(folder, file_name, parsed))
        all_lods.append(build_lod_df(folder, file_name, parsed))
        all_filters.append(build_filters_df(folder, file_name, parsed))
        all_worksheets.append(build_worksheets_df(folder, file_name, parsed))
        all_dashboards.append(build_dashboards_df(folder, file_name, parsed))
        all_actions.append(build_actions_df(folder, file_name, parsed))

    def _concat(frames):
        frames = [f for f in frames if f is not None and not f.empty]
        return pd.concat(frames, ignore_index=True) if frames else pd.DataFrame()

    matched_lookup = defaultdict(list)
    for usage_name, file_name in matched_pairs:
        matched_lookup[usage_name].append(file_name)
    usage_rows = []
    for usage_name in sorted(allowed_usage_names):
        matched_files = sorted(set(matched_lookup.get(usage_name, [])))
        usage_rows.append({
            'Content Usage Entry': usage_name,
            'Matched Workbook File': ', '.join(matched_files) if matched_files else 'Not Found in Uploaded Workbooks',
            'Match Status': 'Matched' if matched_files else 'Not Found',
        })

    table_db_rows, mapped_count, needs_review_count = ([], 0, 0)
    if tds_map:
        table_db_rows, mapped_count, needs_review_count = build_table_db_mapping(unique_tables, tds_map)

    return {
        'overview': pd.DataFrame(overview_rows), 'tables': _concat(all_tables),
        'relationships': _concat(all_rels), 'formulas': _concat(all_formulas), 'lods': _concat(all_lods),
        'filters': _concat(all_filters), 'worksheets': _concat(all_worksheets),
        'dashboards': _concat(all_dashboards), 'actions': _concat(all_actions),
        'duplicates': build_duplicates_df(dup_records), 'usage': pd.DataFrame(usage_rows),
        'validation': pd.DataFrame(validation_rows),
        'table_db_mapping': pd.DataFrame(table_db_rows),
        'table_db_mapped_count': mapped_count, 'table_db_needs_review_count': needs_review_count,
        'unique_tables': unique_tables,
    }


# ===========================================================================
# EXCEL WORKBOOK WRITER (full, write-only for memory safety)
# ===========================================================================

def build_output_workbook(workbook_results, dup_records, allowed_usage_names, matched_pairs, tds_map):
    total = len(workbook_results)
    progress = st.progress(0)
    status = st.empty()
    status.text("Pass 1 of 2: Analyzing parsed workbooks...")

    display_tables = build_display_tables(workbook_results, dup_records, allowed_usage_names, matched_pairs, tds_map)
    counts = duplicate_counts(dup_records)
    progress.progress(0.35)

    status.text("Pass 2 of 2: Creating Excel workbook...")
    wb = Workbook(write_only=True)

    ws_overview = _setup_write_only_sheet(wb, 'Overview',
        ['File Path', 'Workbook Name', 'Datasources', 'Tables', 'Relationships', 'Calculated Fields',
         'Filters', 'LOD Translations', 'Worksheets', 'Dashboards', 'Actions', 'Complexity',
         'Duplicate Type', 'Parse Status'],
        [28, 40, 12, 10, 14, 16, 10, 14, 12, 12, 10, 14, 16, 14])
    ws_tables = _setup_write_only_sheet(wb, 'Tables', TABLES_HEADERS, TABLES_WIDTHS)
    ws_rels = _setup_write_only_sheet(wb, 'Relationships', RELATIONSHIPS_HEADERS, RELATIONSHIPS_WIDTHS)
    ws_formulas = _setup_write_only_sheet(wb, 'Formula Mapping', FORMULA_HEADERS, FORMULA_WIDTHS)
    ws_lods = _setup_write_only_sheet(wb, 'LOD Translations', LOD_HEADERS, LOD_WIDTHS)
    ws_filters = _setup_write_only_sheet(wb, 'Filters', FILTERS_HEADERS, FILTERS_WIDTHS)
    ws_worksheets = _setup_write_only_sheet(wb, 'Worksheets', WORKSHEETS_HEADERS, WORKSHEETS_WIDTHS)
    ws_dashboards = _setup_write_only_sheet(wb, 'Dashboards', DASHBOARDS_HEADERS, DASHBOARDS_WIDTHS)
    ws_actions = _setup_write_only_sheet(wb, 'Actions', ACTIONS_HEADERS, ACTIONS_WIDTHS)
    ws_dup = _setup_write_only_sheet(wb, 'Duplicate Analysis', DUPLICATE_HEADERS, DUPLICATE_WIDTHS)
    ws_usage = _setup_write_only_sheet(wb, 'Content Usage Matches', USAGE_HEADERS, USAGE_WIDTHS)
    ws_table_db = _setup_write_only_sheet(wb, 'Table DB Mapping', TABLE_DB_MAPPING_HEADERS, TABLE_DB_MAPPING_WIDTHS)
    ws_validation = _setup_write_only_sheet(wb, 'Validation Checklist', VALIDATION_HEADERS, VALIDATION_WIDTHS)
    ws_summary = _setup_write_only_sheet(wb, 'Final Summary', FINAL_SUMMARY_HEADERS, FINAL_SUMMARY_WIDTHS)

    for _, row in display_tables['overview'].iterrows():
        _append_row_values(ws_overview, list(row))
    category_colors = {'Exact Duplicate': COLORS['grp_a'], 'Near Duplicate': COLORS['grp_b'],
                        'Same Data Source': COLORS['grp_c'], 'Unique': COLORS['unique'], 'Unparsed': COLORS['unparsed']}
    for _, row in display_tables['tables'].iterrows():
        _append_row_values(ws_tables, list(row), *COLORS['table'])
    for _, row in display_tables['relationships'].iterrows():
        _append_row_values(ws_rels, list(row))
    for _, row in display_tables['formulas'].iterrows():
        _append_row_values(ws_formulas, list(row), *COLORS['calc_measure'])
    for _, row in display_tables['lods'].iterrows():
        _append_row_values(ws_lods, list(row), *COLORS['calc_lod'])
    for _, row in display_tables['filters'].iterrows():
        bg, fg = COLORS['filter_shared'] if str(row.get('Scope', '')).startswith('Shared') else COLORS['filter_worksheet']
        _append_row_values(ws_filters, list(row), bg, fg)
    for _, row in display_tables['worksheets'].iterrows():
        _append_row_values(ws_worksheets, list(row))
    for _, row in display_tables['dashboards'].iterrows():
        _append_row_values(ws_dashboards, list(row))
    for _, row in display_tables['actions'].iterrows():
        _append_row_values(ws_actions, list(row))
    for _, row in display_tables['duplicates'].iterrows():
        bg, fg = category_colors.get(row.get('Duplicate Type', ''), ('FFFFFF', '000000'))
        _append_row_values(ws_dup, list(row), bg, fg, bold=row.get('Duplicate Type') in {'Exact Duplicate', 'Near Duplicate'})
    for _, row in display_tables['usage'].iterrows():
        _append_row_values(ws_usage, list(row))
    for _, row in display_tables['table_db_mapping'].iterrows():
        _append_row_values(ws_table_db, list(row))
    for _, row in display_tables['validation'].iterrows():
        _append_row_values(ws_validation, list(row))

    progress.progress(0.85)

    summary_rows = [
        ('Total Workbooks Processed', total),
        ('Exact Duplicate Workbooks', counts['exact']),
        ('Near Duplicate Workbooks', counts['near']),
        ('Same Data Source Workbooks', counts['same_source']),
        ('Unique Workbooks', counts['unique']),
        ('Unparsed Workbooks', counts['unparsed']),
        ('Total Unique Tables Used', len(display_tables['unique_tables'])),
        ('Total Calculated Fields', int(display_tables['overview']['Calculated Fields'].sum()) if 'Calculated Fields' in display_tables['overview'] else 0),
        ('Total LOD Translations', int(display_tables['overview']['LOD Translations'].sum()) if 'LOD Translations' in display_tables['overview'] else 0),
        ('Total Filters', int(display_tables['overview']['Filters'].sum()) if 'Filters' in display_tables['overview'] else 0),
        ('Total Worksheets', int(display_tables['overview']['Worksheets'].sum()) if 'Worksheets' in display_tables['overview'] else 0),
        ('Total Dashboards', int(display_tables['overview']['Dashboards'].sum()) if 'Dashboards' in display_tables['overview'] else 0),
        ('Table DB Mapped', display_tables['table_db_mapped_count']),
        ('Table DB Needs Review', display_tables['table_db_needs_review_count']),
    ]
    for row_values in summary_rows:
        _append_row_values(ws_summary, row_values)

    progress.progress(0.98)
    status.text("Saving Excel workbook...")
    output = BytesIO()
    wb.save(output)
    output.seek(0)
    progress.progress(1.0)
    status.text("Excel workbook ready.")

    return output, display_tables, counts


def build_single_workbook_excel(workbook_row, tables_df, rels_df, formulas_df, lods_df, filters_df,
                                 worksheets_df, dashboards_df, actions_df):
    """Builds a small multi-sheet Excel workbook containing only the selected
    workbook's data, mirroring build_single_report_excel() from the
    WebFOCUS version."""
    output = BytesIO()
    with pd.ExcelWriter(output, engine='openpyxl') as writer:
        pd.DataFrame([workbook_row]).astype(str).to_excel(writer, sheet_name='Overview', index=False)
        sheets = [
            ('Tables', tables_df), ('Relationships', rels_df), ('Formulas', formulas_df),
            ('LOD Translations', lods_df), ('Filters', filters_df), ('Worksheets', worksheets_df),
            ('Dashboards', dashboards_df), ('Actions', actions_df),
        ]
        for sheet_name, df in sheets:
            if df is not None and not df.empty:
                df.astype(str).to_excel(writer, sheet_name=sheet_name[:31], index=False)
    output.seek(0)
    return output


# ===========================================================================
# DUPLICATE CODE DIFF VIEWER (compares raw .twb XML text side by side)
# ===========================================================================

def compute_diff_blocks(text_a, text_b):
    lines_a = text_a.splitlines()
    lines_b = text_b.splitlines()
    sm = difflib.SequenceMatcher(None, lines_a, lines_b, autojunk=False)
    blocks = []

    def numbered(lines, start_idx):
        return [(start_idx + offset + 1, line) for offset, line in enumerate(lines)]

    for tag, i1, i2, j1, j2 in sm.get_opcodes():
        blocks.append((tag, numbered(lines_a[i1:i2], i1), numbered(lines_b[j1:j2], j1)))
    return blocks


def diff_summary_counts(blocks):
    removed = added = changed = 0
    for tag, a, b in blocks:
        if tag == 'delete':
            removed += len(a)
        elif tag == 'insert':
            added += len(b)
        elif tag == 'replace':
            changed += max(len(a), len(b))
    return removed, added, changed


def render_diff_blocks_html(blocks, name_a, name_b):
    rows_html = []
    for tag, a, b in blocks:
        if tag == 'equal':
            for (num_a, line_a), (num_b, line_b) in zip(a, b):
                rows_html.append(
                    f'<tr><td class="diff-linenum">{num_a}</td><td class="diff-cell-equal">{escape(line_a)}</td>'
                    f'<td class="diff-linenum">{num_b}</td><td class="diff-cell-equal">{escape(line_b)}</td></tr>'
                )
        elif tag == 'delete':
            for num_a, line_a in a:
                rows_html.append(
                    f'<tr><td class="diff-linenum">{num_a}</td><td class="diff-cell-removed">{escape(line_a)}</td>'
                    f'<td class="diff-linenum"></td><td class="diff-cell-empty"></td></tr>'
                )
        elif tag == 'insert':
            for num_b, line_b in b:
                rows_html.append(
                    f'<tr><td class="diff-linenum"></td><td class="diff-cell-empty"></td>'
                    f'<td class="diff-linenum">{num_b}</td><td class="diff-cell-added">{escape(line_b)}</td></tr>'
                )
        elif tag == 'replace':
            max_len = max(len(a), len(b))
            for idx in range(max_len):
                if idx < len(a):
                    num_a, line_a = a[idx]
                    left = f'<td class="diff-linenum">{num_a}</td><td class="diff-cell-removed">{escape(line_a)}</td>'
                else:
                    left = '<td class="diff-linenum"></td><td class="diff-cell-empty"></td>'
                if idx < len(b):
                    num_b, line_b = b[idx]
                    right = f'<td class="diff-linenum">{num_b}</td><td class="diff-cell-added">{escape(line_b)}</td>'
                else:
                    right = '<td class="diff-linenum"></td><td class="diff-cell-empty"></td>'
                rows_html.append(f'<tr>{left}{right}</tr>')

    return f"""
    <div class="kash-table-wrap" style="max-height: 520px;">
        <table class="diff-side-by-side">
            <thead><tr>
                <th class="diff-linenum-header">#</th><th>{escape(name_a)}</th>
                <th class="diff-linenum-header">#</th><th>{escape(name_b)}</th>
            </tr></thead>
            <tbody>{''.join(rows_html)}</tbody>
        </table>
    </div>
    """


def find_duplicate_match_options(duplicate_df, overview_df, selected_name, selected_path):
    if duplicate_df is None or duplicate_df.empty:
        return []
    row_matches = duplicate_df[(duplicate_df['Workbook Name'] == selected_name) & (duplicate_df['File Path'] == selected_path)]
    if row_matches.empty:
        return []
    matched_names_raw = str(row_matches.iloc[0].get('Matched With', '') or '')
    matched_names = [n.strip() for n in matched_names_raw.split(',') if n.strip()]
    options = []
    for matched_name in matched_names:
        candidates = overview_df[overview_df['Workbook Name'] == matched_name]
        for _, cand_row in candidates.iterrows():
            cand_path = cand_row['File Path']
            if matched_name == selected_name and cand_path == selected_path:
                continue
            options.append((f"{matched_name} | {cand_path}", matched_name, cand_path))
    seen, deduped = set(), []
    for label, name, path in options:
        key = (name, path)
        if key not in seen:
            seen.add(key)
            deduped.append((label, name, path))
    return deduped


def render_duplicate_code_diff(workbook_row, duplicate_df, overview_df, raw_text_lookup, selected_name, selected_path):
    dup_type = str(workbook_row.get('Duplicate Type', '') or '')
    if dup_type not in {'Exact Duplicate', 'Near Duplicate'}:
        return

    st.markdown("---")
    st.subheader("Compare XML With Matched Duplicate")
    match_options = find_duplicate_match_options(duplicate_df, overview_df, selected_name, selected_path)
    if not match_options:
        st.info("This workbook is flagged as a duplicate, but no comparable matched file could be located.")
        return

    labels = [label for label, _, _ in match_options]
    chosen_label = st.selectbox(
        "Compare against", labels,
        key=f"dup_compare_choice_{re.sub(r'[^A-Za-z0-9_]+', '_', selected_name + selected_path)}",
    )
    chosen = next((m for m in match_options if m[0] == chosen_label), None)
    if not chosen:
        return
    _, match_name, match_path = chosen

    text_a = raw_text_lookup.get((selected_path, selected_name), '')
    text_b = raw_text_lookup.get((match_path, match_name), '')
    if not text_a or not text_b:
        st.warning("Raw workbook XML for one or both files is not available for comparison.")
        return
    if text_a == text_b:
        st.success(f"'{selected_name}' and '{match_name}' are byte-for-byte identical.")
        return

    blocks = compute_diff_blocks(text_a, text_b)
    removed, added, changed = diff_summary_counts(blocks)
    if removed == 0 and added == 0 and changed == 0:
        st.success(f"'{selected_name}' and '{match_name}' have no line-level differences.")
        return

    c1, c2, c3 = st.columns(3)
    c1.metric(f"Only in {selected_name}", removed)
    c2.metric(f"Only in {match_name}", added)
    c3.metric("Changed line pairs", changed)
    st.caption("🔴 Red = only in the currently selected workbook. 🟢 Green = only in the compared file. White = identical in both.")
    st.markdown(render_diff_blocks_html(blocks, selected_name, match_name), unsafe_allow_html=True)


# ===========================================================================
# BADGES + SMALL RENDER HELPERS
# ===========================================================================

def badge(label, level=''):
    css_class = 'badge'
    if level:
        css_class += f' badge-{level}'
    return f'<span class="{css_class}">{label}</span>'


def show_table_or_info(df, message, large=False, height="content"):
    if df is None or df.empty:
        st.info(message)
        return
    if large or len(df) > 150:
        st.dataframe(df, use_container_width=True, hide_index=True, height=height)
        return
    html = df.to_html(index=False, escape=True, classes="kash-table", border=0)
    st.markdown(f'<div class="kash-table-wrap">{html}</div>', unsafe_allow_html=True)


def _render_plan_step(step):
    level_map = {'confirmed': 'ok', 'rule_based': 'ok', 'inferred': 'medium', 'suggested': 'medium', 'manual_review': 'high'}
    level = level_map.get(step.classification, 'medium')
    with st.container(border=True):
        st.markdown(f"**{step.title}** — {step.power_bi_layer}")
        st.markdown(
            f"{badge(step.classification.replace('_', ' ').title(), level)} {badge(f'Confidence: {step.confidence}%')}",
            unsafe_allow_html=True,
        )
        st.write(step.action)
        st.caption(f"Why: {step.reason}")
        if step.referenced_sources or step.referenced_fields:
            st.caption(f"References: {', '.join(step.referenced_sources + step.referenced_fields)}")
        if step.manual_validation:
            st.warning(step.manual_validation)


def render_ai_build_plan(workbook_row, parsed):
    st.markdown(
        "Generates a **hybrid rule-based + AI** build plan: LOD expressions are translated "
        "deterministically (no AI, no cost), and one AI call organizes the remaining judgment "
        "calls - relationship cardinality risk, table calculations, novel formulas, and page/visual suggestions."
    )

    st.markdown("#### Rule-Based Translations (no AI, zero cost)")
    if not parsed['lod_translations']:
        st.info("No LOD (FIXED/INCLUDE/EXCLUDE) expressions were detected in this workbook.")
    for lod in parsed['lod_translations']:
        with st.container(border=True):
            st.markdown(f"**{lod['field']}** ({lod['lod_type']} on {', '.join(lod['dims']) or '(no dims)'})")
            st.code(lod['dax'], language='DAX')
            if lod['caveat']:
                st.warning(lod['caveat'])

    st.markdown("---")

    client = get_openai_client()
    if client is None:
        st.warning(
            "OpenAI is not configured. Add `OPENAI_API_KEY` to `.streamlit/secrets.toml` "
            "(e.g. `OPENAI_API_KEY = \"sk-...\"`) to enable the AI-assisted portion."
        )
        return

    workbook_key = re.sub(r'[^A-Za-z0-9_]+', '_', str(workbook_row.get('Workbook Name', '')) + '_' + str(workbook_row.get('File Path', '')))
    cache = st.session_state.setdefault('ai_build_plan_cache', {})

    generate_clicked = st.button(
        "Generate AI build plan" if workbook_key not in cache else "Regenerate AI build plan",
        key=f"ai_build_plan_btn_{workbook_key}",
    )

    if generate_clicked:
        payload = build_ai_build_plan_payload(workbook_row, parsed)
        with st.spinner("Calling OpenAI (one request for this workbook)..."):
            plan, error = generate_ai_build_plan(client, "gpt-4.1-mini", payload)
        cache[workbook_key] = (plan, error)

    cached = cache.get(workbook_key)
    if cached is None:
        st.info("Click the button above to generate the AI-assisted portion of the build plan.")
        return

    plan, error = cached
    if error:
        st.error(error)
        return

    summary = plan.report_summary
    st.markdown("#### AI-Assisted Plan")
    c1, c2, c3 = st.columns(3)
    c1.metric("Complexity", summary.complexity_level)
    c2.metric("Overall confidence", f"{summary.overall_confidence}%")
    c3.metric("Manual review items", len(plan.manual_review_items))
    st.write(f"**Likely purpose:** {summary.likely_purpose}")
    st.write(f"**Report type:** {summary.report_type}")

    tabs = st.tabs(["Implementation Steps", "Data Prep", "Calculations", "Filters/Relationships", "Pages", "Validation", "Manual Review"])
    with tabs[0]:
        if not plan.implementation_plan:
            st.info("No implementation steps returned.")
        for step in sorted(plan.implementation_plan, key=lambda s: s.step_number):
            _render_plan_step(step)
    with tabs[1]:
        if not plan.data_preparation:
            st.info("No data preparation steps returned.")
        for step in plan.data_preparation:
            _render_plan_step(step)
    with tabs[2]:
        if not plan.calculation_plan:
            st.info("No calculation translations returned.")
        for step in plan.calculation_plan:
            _render_plan_step(step)
    with tabs[3]:
        if not plan.filter_parameter_plan:
            st.info("No filter/relationship items returned.")
        for step in plan.filter_parameter_plan:
            _render_plan_step(step)
    with tabs[4]:
        if not plan.page_recommendations:
            st.info("No page/visual recommendations returned.")
        for step in plan.page_recommendations:
            _render_plan_step(step)
    with tabs[5]:
        if not plan.validation_checks:
            st.info("No validation checks returned.")
        for idx, check in enumerate(plan.validation_checks, start=1):
            st.write(f"{idx}. {check}")
    with tabs[6]:
        if not plan.manual_review_items:
            st.success("No manual review items flagged.")
        for idx, item in enumerate(plan.manual_review_items, start=1):
            st.warning(f"{idx}. {item}")
        if plan.limitations:
            st.caption("Limitations:")
            for lim in plan.limitations:
                st.caption(f"• {lim}")


# ===========================================================================
# "START HERE" NATURAL-LANGUAGE SUMMARY
# ===========================================================================

def unique_text_values(values):
    seen, result = set(), []
    for value in values or []:
        text = str(value or '').strip()
        if not text or text.lower() in {'nan', 'none'}:
            continue
        key = text.upper()
        if key not in seen:
            seen.add(key)
            result.append(text)
    return result


def build_start_here_facts(workbook_row, parsed):
    sources = []
    for ds in parsed['datasources']:
        sources.extend(t['friendly_name'] for t in ds['tables'])
    sources = unique_text_values(sources)

    return {
        'workbook': str(workbook_row.get('Workbook Name', 'Selected workbook')),
        'sources': sources,
        'relationship_count': len(parsed['relationships']),
        'formula_count': len(parsed['formula_mapping']),
        'lod_count': len(parsed['lod_translations']),
        'filter_count': len(parsed['filters']),
        'worksheet_count': len(parsed['worksheets']),
        'dashboard_count': len(parsed['dashboards']),
        'action_count': len(parsed['actions']),
        'duplicate_status': str(workbook_row.get('Duplicate Type', 'Unique') or 'Unique'),
        'complexity': str(workbook_row.get('Complexity', 'Medium') or 'Medium'),
    }


def fallback_workbook_summary(facts):
    source_count = len(facts['sources'])
    source_word = 'source' if source_count == 1 else 'sources'
    rel_word = 'relationship' if facts['relationship_count'] == 1 else 'relationships'
    formula_word = 'calculated field' if facts['formula_count'] == 1 else 'calculated fields'
    filter_word = 'filter' if facts['filter_count'] == 1 else 'filters'
    ws_word = 'worksheet' if facts['worksheet_count'] == 1 else 'worksheets'
    source_text = ', '.join(facts['sources']) if facts['sources'] else 'the detected Tableau sources'

    duplicate_status = facts['duplicate_status']
    if duplicate_status in {'Exact Duplicate', 'Near Duplicate', 'Same Data Source'}:
        duplicate_text = f"{duplicate_status} - rebuild once where possible and reuse the same Power BI approach for this group."
    else:
        duplicate_text = "Unique - no duplicate group was detected for this workbook."

    purpose = f"Reads {source_text} and builds {facts['worksheet_count']} {ws_word}"
    if facts['dashboard_count']:
        purpose += f" combined into {facts['dashboard_count']} dashboard(s)"
    purpose += "."

    complexity_text = (
        f"{facts['complexity']} - {source_count} {source_word}, {facts['relationship_count']} {rel_word}, "
        f"{facts['formula_count']} {formula_word} ({facts['lod_count']} using Level of Detail expressions), "
        f"and {facts['filter_count']} {filter_word}."
    )

    if facts['relationship_count']:
        direction = (
            "Confirm relationship cardinality first (Tableau relationships don't declare it), then build the "
            "data model and Power Query merge steps, followed by calculations and visuals."
        )
    else:
        direction = "Load the source table(s) into Power Query, then build calculations and visuals directly."

    return {'purpose': purpose, 'complexity': complexity_text, 'duplicate_status': duplicate_text, 'recommended_direction': direction}


def render_start_here(workbook_row, parsed):
    facts = build_start_here_facts(workbook_row, parsed)
    summary = fallback_workbook_summary(facts)
    source_text = ', '.join(facts['sources']) if facts['sources'] else 'No source tables detected'

    with st.container(border=True):
        st.markdown(f"**Workbook:** {escape(facts['workbook'])}")
        st.markdown(f"**Purpose:** {escape(summary['purpose'])}")
        st.markdown(f"**Main sources:** {escape(source_text)}")
        st.markdown(f"**Migration complexity:** {escape(summary['complexity'])}")
        st.markdown(f"**Duplicate status:** {escape(summary['duplicate_status'])}")
        st.markdown(f"**Recommended direction:** {escape(summary['recommended_direction'])}")


def collect_review_items(parsed):
    items = []
    for r in parsed['relationships']:
        items.append(('High', 'Relationship cardinality check',
                       f"{r['left_field']} {r['join_type']} {r['right_field']} ({r['style']})",
                       'Tableau relationships do not declare cardinality - confirm key uniqueness before building the Power BI model relationship.'))
    for lod in parsed['lod_translations']:
        if lod['caveat']:
            items.append(('Medium', f"LOD: {lod['field']}", lod['agg_expr_raw'], lod['caveat']))
    for f in parsed['formula_mapping']:
        if 'Table calculation' in f['formula_type']:
            items.append(('High', f"Table calculation: {f['field']}", f['formula_raw'],
                          'Depends on the original view\'s partition/addressing - confirm against the source worksheet before finalizing the DAX.'))
    for filt in parsed['filters']:
        if filt['filter_class'] == 'measure-names':
            items.append(('Medium', 'Measure Names filter', ', '.join(filt['measures_shown']),
                          'Rebuild as a Power BI Field Parameter, not a slicer on a real column.'))
    return items


def render_manual_review(parsed):
    items = collect_review_items(parsed)
    if not items:
        st.success("No high-priority manual review items were detected.")
        return
    priority_order = {'High': 0, 'Medium': 1, 'Low': 2}
    for idx, (priority, title, reason, action) in enumerate(sorted(items, key=lambda x: priority_order.get(x[0], 9)), start=1):
        level = priority.lower()
        with st.expander(f"{idx}. {priority} - {title}", expanded=idx <= 5):
            st.markdown(badge(priority, level if level in {'high', 'medium', 'low'} else ''), unsafe_allow_html=True)
            st.write(f"Detail: {reason}")
            st.write(f"Action: {action}")


# ===========================================================================
# STREAMLIT PAGE LAYOUT
# ===========================================================================

st.set_page_config(page_title="KashMap Migration Workspace", layout="wide")

st.markdown(
    """
<style>
:root {
    --kash-bg: var(--background-color, #FFFFFF);
    --kash-section: var(--secondary-background-color, #FBFAF8);
    --kash-card: var(--background-color, #FFFFFF);
    --kash-navy: var(--text-color, #1F2933);
    --kash-deep: var(--text-color, #111827);
    --kash-blue: #2563EB;
    --kash-blue-soft: var(--secondary-background-color, #EAF2FF);
    --kash-gold: #FF5A1F;
    --kash-green: #2E7D32;
    --kash-orange: #FF5A1F;
    --kash-orange-dark: #E94D13;
    --kash-orange-soft: var(--secondary-background-color, #FFF3EC);
    --kash-red: #B42318;
    --kash-border: color-mix(in srgb, var(--text-color, #1F2933) 18%, transparent);
    --kash-muted: color-mix(in srgb, var(--text-color, #1F2933) 55%, transparent);
}

.upload-card, .upload-card-small,
.wizard-card, .workspace-card,
.report-header, .kpi-strip {
    background: rgba(127, 127, 127, 0.08) !important;
    border-color: var(--kash-border) !important;
}

.main .block-container {
    padding-top: 1.4rem;
    padding-bottom: 3rem;
    max-width: 1480px;
}
.kash-hero {
    position: relative;
    overflow: hidden;
    isolation: isolate;
    border: 1px solid #FFE0D2;
    background: linear-gradient(135deg, #FFFFFF 0%, #FFFFFF 58%, #FFF7F2 100%);
    border-radius: 8px;
    padding: 30px 34px;
    margin-bottom: 20px;
    box-shadow: 0 18px 44px rgba(16, 24, 40, 0.08);
}
.kash-hero::before {
    content: "";
    position: absolute;
    top: 0;
    right: 0;
    bottom: 0;
    width: 56%;
    z-index: 0;
    opacity: 0.72;
    pointer-events: none;
    background-image: url("data:image/svg+xml,%3Csvg width='760' height='220' viewBox='0 0 760 220' xmlns='http://www.w3.org/2000/svg'%3E%3Cg fill='none' stroke='%23ff5a1f' stroke-width='1.35' stroke-linecap='round'%3E%3Cpath opacity='.28' d='M18 178 C110 40 226 34 330 104 S520 184 732 42'/%3E%3Cpath opacity='.18' d='M52 22 C142 116 238 134 372 72 S598 8 736 132'/%3E%3Cpath opacity='.22' d='M112 212 C204 92 310 82 418 142 S586 202 750 94'/%3E%3Cpath opacity='.16' d='M0 92 L126 46 L238 128 L364 56 L498 144 L646 84 L760 124'/%3E%3Cpath opacity='.14' d='M84 168 L210 84 L338 162 L466 82 L608 152 L724 50'/%3E%3C/g%3E%3Cg fill='%23111827'%3E%3Ccircle opacity='.45' cx='126' cy='46' r='3'/%3E%3Ccircle opacity='.35' cx='238' cy='128' r='2.5'/%3E%3Ccircle opacity='.4' cx='364' cy='56' r='3'/%3E%3Ccircle opacity='.35' cx='498' cy='144' r='2.5'/%3E%3Ccircle opacity='.42' cx='646' cy='84' r='3'/%3E%3C/g%3E%3Cg fill='%23ff5a1f'%3E%3Ccircle opacity='.26' cx='210' cy='84' r='4'/%3E%3Ccircle opacity='.24' cx='418' cy='142' r='4'/%3E%3Ccircle opacity='.22' cx='608' cy='152' r='4'/%3E%3C/g%3E%3C/svg%3E");
    background-repeat: no-repeat;
    background-size: cover;
    background-position: right center;
}
.kash-title {
    position: relative;
    z-index: 1;
    font-size: 2.8rem;
    font-weight: 900;
    color: var(--kash-orange);
    margin: 0;
    letter-spacing: 0;
}
.kash-title-kash { color: #111827; }
.kash-title-map { color: var(--kash-orange); }
.kash-subtitle {
    position: relative;
    z-index: 1;
    font-size: 1.15rem;
    font-weight: 700;
    color: #111827;
    margin-top: 8px;
    max-width: 980px;
    line-height: 1.55;
}
.kash-subtitle-accent { color: var(--kash-orange); }
.section-label {
    font-size: 1.2rem;
    font-weight: 800;
    margin: 18px 0 8px 0;
}
.upload-card {
    background: rgba(127, 127, 127, 0.08);
    border: 1px solid var(--kash-border);
    border-radius: 16px;
    padding: 20px 22px;
    box-shadow: 0 1px 3px rgba(16, 24, 40, 0.06);
    min-height: 100%;
}
.upload-card-small {
    background: rgba(127, 127, 127, 0.08);
    border: 1px solid var(--kash-border);
    border-radius: 16px;
    padding: 18px 20px;
    box-shadow: 0 1px 3px rgba(16, 24, 40, 0.06);
    margin-bottom: 16px;
}
.kash-table-wrap {
    border: 1px solid var(--kash-border);
    border-radius: 14px;
    overflow: auto;
    max-height: 640px;
    margin: 14px 0 22px 0;
    box-shadow: 0 8px 24px rgba(17, 24, 39, 0.08);
    background: rgba(127, 127, 127, 0.08);
}
.diff-side-by-side {
    width: 100%;
    border-collapse: collapse;
    font-family: 'Consolas', 'Menlo', monospace;
    font-size: 0.82rem;
    table-layout: fixed;
}
.diff-side-by-side th {
    background: #1F2933;
    color: #FFFFFF;
    text-align: left;
    padding: 8px 10px;
    font-family: Arial, sans-serif;
    font-size: 0.82rem;
    position: sticky;
    top: 0;
}
.diff-side-by-side th.diff-linenum-header { width: 48px; }
.diff-side-by-side td {
    padding: 3px 10px;
    vertical-align: top;
    white-space: pre-wrap;
    word-break: break-word;
    border-bottom: 1px solid #EEF1F5;
}
.diff-side-by-side td.diff-linenum {
    width: 48px;
    text-align: right;
    color: #94A3B8;
    background: #FAFBFC;
    user-select: none;
    padding-right: 8px;
    border-right: 1px solid #EEF1F5;
    white-space: nowrap;
}
.diff-cell-removed { background: #FFEBEE; color: #B42318; }
.diff-cell-added { background: #E8F5E9; color: #1B5E20; }
.diff-cell-equal { background: #FFFFFF; color: #444B56; }
.diff-cell-empty { background: #FAFAFA; }
.kash-table thead th { position: sticky; top: 0; z-index: 1; }
.kash-table {
    width: 100%;
    border-collapse: collapse;
    font-size: 0.9rem;
}
.kash-table thead th {
    background: linear-gradient(180deg, #2B2B2B 0%, #1F2933 100%);
    color: #FFFFFF;
    font-weight: 800;
    text-align: left;
    padding: 13px 15px;
    border-right: 1px solid rgba(255, 255, 255, 0.12);
    border-bottom: 3px solid #FF5A1F;
    white-space: nowrap;
    letter-spacing: 0.01em;
}
.kash-table thead th:last-child { border-right: none; }
.kash-table tbody td {
    padding: 12px 15px;
    border-bottom: 1px solid rgba(127, 127, 127, 0.2);
    border-right: 1px solid rgba(127, 127, 127, 0.12);
    vertical-align: top;
}
.kash-table tbody td:last-child { border-right: none; }
.kash-table tbody tr:nth-child(even) { background: rgba(127, 127, 127, 0.06); }
.kash-table tbody tr:hover { background: rgba(255, 90, 31, 0.12); }
.kash-table tbody tr:last-child td { border-bottom: none; }
.kash-table td:first-child { font-weight: 700; }
div[data-testid="stDataFrame"] {
    border: 1px solid var(--kash-border);
    border-radius: 14px;
    overflow: hidden;
    box-shadow: 0 4px 14px rgba(16, 24, 40, 0.06);
    margin: 14px 0 22px 0;
}
.card-title { font-size: 1rem; font-weight: 800; margin-bottom: 6px; }
.card-subtitle { font-size: 0.86rem; opacity: 0.7; margin-bottom: 14px; }
.badge-required {
    width: fit-content;
    margin-left: auto;
    font-size: 0.72rem;
    font-weight: 700;
    color: #B42318;
    background: #FFF1F0;
    border: 1px solid #FFE1DF;
    border-radius: 999px;
    padding: 6px 10px;
    text-align: center;
}
.badge-optional {
    width: fit-content;
    margin-left: auto;
    font-size: 0.72rem;
    font-weight: 700;
    opacity: 0.75;
    background: rgba(127, 127, 127, 0.12);
    border: 1px solid var(--kash-border);
    border-radius: 999px;
    padding: 6px 10px;
    text-align: center;
}
.file-count-box {
    background: rgba(127, 127, 127, 0.08);
    border: 1px solid var(--kash-border);
    border-radius: 12px;
    padding: 10px 12px;
    margin-top: 10px;
    font-size: 0.88rem;
    font-weight: 600;
}
.card-title-row { display: flex; align-items: center; justify-content: space-between; gap: 12px; margin-bottom: 4px; }
.wizard-card {
    background: rgba(127, 127, 127, 0.08);
    border: 1px solid var(--kash-border);
    border-radius: 8px;
    padding: 22px;
    margin-bottom: 16px;
    box-shadow: 0 8px 24px rgba(15, 23, 42, 0.06);
}
.wizard-card h3 { font-size: 1.15rem; margin: 0 0 10px 0; }
.wizard-card p { line-height: 1.45; margin: 4px 0 10px 0; }
div[data-testid="stExpander"]:hover { border-left: 4px solid var(--kash-orange); }
.badge {
    display: inline-block;
    padding: 5px 10px;
    border-radius: 999px;
    background-color: var(--kash-blue-soft);
    color: var(--kash-navy);
    font-size: 0.78rem;
    font-weight: 700;
    margin-right: 6px;
    margin-bottom: 6px;
}
.badge-high { background-color: var(--kash-navy); color: #FFFFFF; }
.badge-medium { background-color: #FEF3C7; color: #92400E; }
.badge-low { background-color: #DCFCE7; color: #166534; }
.badge-ok { background-color: #DCFCE7; color: #166534; }
.step-title { font-weight: 800; margin-bottom: 4px; }
.workspace-card {
    background: rgba(127, 127, 127, 0.08);
    border: 1px solid var(--kash-border);
    border-radius: 8px;
    padding: 20px;
    box-shadow: 0 8px 24px rgba(15, 23, 42, 0.055);
    margin-bottom: 14px;
}
.report-header {
    background: rgba(127, 127, 127, 0.08);
    border: 1px solid var(--kash-border);
    border-radius: 8px;
    padding: 22px 24px;
    box-shadow: 0 10px 28px rgba(15, 23, 42, 0.06);
    margin: 12px 0 16px 0;
}
.kpi-strip {
    background: rgba(127, 127, 127, 0.08);
    border: 1px solid var(--kash-border);
    border-radius: 8px;
    padding: 16px;
    box-shadow: 0 8px 20px rgba(15, 23, 42, 0.045);
}
.stButton > button, .stDownloadButton > button {
    border-radius: 10px;
    border: 1px solid var(--kash-border);
    font-weight: 700;
}
.stButton > button[kind="primary"] {
    background-color: #FF5A1F !important;
    border-color: #FF5A1F !important;
    color: white !important;
}
.stButton > button[kind="primary"]:hover {
    background-color: var(--kash-orange-dark) !important;
    border-color: var(--kash-orange-dark) !important;
    color: white !important;
}
@media (max-width: 700px) {
    .kash-title { font-size: 2rem; }
    .kash-hero { padding: 24px; }
}
</style>
    """,
    unsafe_allow_html=True,
)

st.markdown(
    """
    <div class="kash-hero">
      <div class="kash-title"><span class="kash-title-kash">Kash</span><span class="kash-title-map">Map</span></div>
      <div class="kash-subtitle">Break Down <span class="kash-subtitle-accent">Tableau</span> Complexity for Faster <span class="kash-subtitle-accent">Power BI</span> Migration</div>
    </div>
    """,
    unsafe_allow_html=True,
)

with st.expander("How to use KashMap", expanded=True):
    st.markdown(
        """
        1. Upload one or more `.twb` / `.twbx` files, or a ZIP folder containing workbooks.
        2. Optional: upload a Content Usage export to match or filter active workbooks.
        3. Optional: upload a TDS/TDSX metadata ZIP to map tables to actual database tables.
        4. Click **Analyze Workbooks**.
        5. Review the overview first, then open the guided Power BI build plan (including the AI-assisted plan).
        6. Use the raw tables only when you need technical detail.
        """
    )

st.markdown('<div class="section-label">Upload Files</div>', unsafe_allow_html=True)

upload_col, side_col = st.columns([1, 1], gap="large")

with upload_col:
    upload_card = st.container(border=True)
    with upload_card:
        title_col, badge_col = st.columns([0.82, 0.18])
        with title_col:
            st.markdown(
                """
                <div class="card-title">Upload Tableau Workbooks</div>
                <div class="card-subtitle">Upload multiple .twb/.twbx files or one ZIP folder containing workbooks.</div>
                """,
                unsafe_allow_html=True,
            )
        with badge_col:
            st.markdown('<div class="badge-required">Required</div>', unsafe_allow_html=True)

        mode = st.radio("Workbook input type", ["Individual Workbook Files", "ZIP File"], horizontal=True, key="twb_input_type")

        if mode == "Individual Workbook Files":
            uploaded_twb_files = st.file_uploader(
                "Upload one or more .twb or .twbx files", type=["twb", "twbx"],
                accept_multiple_files=True, key="twb_files_uploader",
            )
            uploaded_zip = None
            if uploaded_twb_files:
                total_size = sum(getattr(f, "size", 0) for f in uploaded_twb_files)
                st.markdown(
                    f'<div class="file-count-box">{len(uploaded_twb_files)} workbook file(s) uploaded · {total_size / 1024:.1f} KB total</div>',
                    unsafe_allow_html=True,
                )
        else:
            uploaded_zip = st.file_uploader("Upload ZIP file containing .twb/.twbx workbooks", type=["zip"], key="zip_file_uploader")
            uploaded_twb_files = []
            if uploaded_zip:
                st.markdown(
                    f'<div class="file-count-box">ZIP uploaded · {uploaded_zip.size / 1024:.1f} KB</div>',
                    unsafe_allow_html=True,
                )

with side_col:
    usage_card = st.container(border=True)
    with usage_card:
        title_col, badge_col = st.columns([0.78, 0.22])
        with title_col:
            st.markdown(
                """
                <div class="card-title">Content Usage File</div>
                <div class="card-subtitle">Upload this only if you want to match or filter active workbooks (Tableau Server/Cloud usage export).</div>
                """,
                unsafe_allow_html=True,
            )
        with badge_col:
            st.markdown('<div class="badge-optional">Optional</div>', unsafe_allow_html=True)
        usage_file = st.file_uploader(
            "Upload Content Usage file", type=["xlsx", "xls", "csv"],
            help="Optional. Upload this only if you want to match/filter workbooks using a Tableau Server/Cloud content usage export.",
            key="usage_uploader",
        )

    st.markdown("<div style='height: 10px;'></div>", unsafe_allow_html=True)

    metadata_card = st.container(border=True)
    with metadata_card:
        title_col, badge_col = st.columns([0.78, 0.22])
        with title_col:
            st.markdown(
                """
                <div class="card-title">TDS/TDSX Metadata ZIP</div>
                <div class="card-subtitle">Optional. Upload a ZIP of published Tableau Data Source files to map tables to actual database tables.</div>
                """,
                unsafe_allow_html=True,
            )
        with badge_col:
            st.markdown('<div class="badge-optional">Optional</div>', unsafe_allow_html=True)
        metadata_zip_file = st.file_uploader(
            "Upload TDS/TDSX metadata ZIP", type=["zip"],
            help="ZIP containing your published .tds / .tdsx data source files.",
            key="metadata_zip_uploader",
        )
        if metadata_zip_file:
            st.markdown(
                f'<div class="file-count-box">Metadata ZIP uploaded · {metadata_zip_file.size / 1024:.1f} KB</div>',
                unsafe_allow_html=True,
            )

    st.markdown("<div style='height: 10px;'></div>", unsafe_allow_html=True)

    output_card = st.container(border=True)
    with output_card:
        st.markdown(
            """
            <div class="card-title">Output Settings</div>
            <div class="card-subtitle">Name of the generated Excel migration workbook.</div>
            """,
            unsafe_allow_html=True,
        )
        output_name = st.text_input("Output file name", value="KashMap_Tableau_Migration_Output.xlsx", key="output_file_name")

st.markdown('<div class="section-label">Manual Expression Translator</div>', unsafe_allow_html=True)
with st.expander("Translate one Tableau expression or query", expanded=False):
    manual_type = st.selectbox(
        "Expression type",
        ["Calculated Field / LOD", "Filter", "Relationship / Join", "Other"],
        key="manual_llm_type",
    )
    manual_context = st.text_input(
        "Optional context", placeholder="Example: field name, source table, workbook name, expected output",
        key="manual_llm_context",
    )
    manual_expression = st.text_area("Paste Tableau expression/formula", height=150, key="manual_llm_expression")
    if st.button("Translate to Power BI", key="manual_llm_translate"):
        client = get_openai_client()
        if client is None:
            st.error("OpenAI is not configured. Add OPENAI_API_KEY to .streamlit/secrets.toml and make sure openai is in requirements.txt.")
        elif not manual_expression.strip():
            st.warning("Paste a Tableau expression or formula first.")
        else:
            with st.spinner("Translating with LLM..."):
                result = translate_expression_with_llm(client, "gpt-4.1-mini", manual_type, manual_expression, manual_context)
            st.subheader("Power BI Translation")
            st.write("**Suggested Power BI Type**")
            st.write(result.get("suggested_power_bi_type", "") or "Not returned")
            st.write("**Suggested Power BI Expression / Action**")
            st.code(result.get("suggested_power_bi_expression_action", "") or "Not returned", language="DAX")
            st.write("**LLM Notes**")
            st.write(result.get("llm_notes", "") or "Not returned")
            st.write("**Needs Manual Review**")
            st.write(result.get("needs_manual_review", "") or "Yes")
            st.write("**Review Reason**")
            st.write(result.get("review_reason", "") or "Review before using in Power BI.")
            st.write("**Confidence**")
            st.write(result.get("confidence", "") or "Not returned")

submit = st.button("Analyze Workbooks", type="primary", use_container_width=True)


@st.fragment
def show_download_button(output_stream, file_name, key):
    st.download_button(
        label="Download Full Excel Report", data=output_stream, file_name=file_name,
        mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        key=key, use_container_width=True,
    )


def filter_report_df(df, workbook_name, file_path):
    if df is None or df.empty or 'Workbook Name' not in df.columns or 'File Path' not in df.columns:
        return pd.DataFrame()
    return df[(df['Workbook Name'] == workbook_name) & (df['File Path'] == file_path)]


def sort_overview_df(df, sort_by):
    if df is None or df.empty:
        return df
    if sort_by == "Name (Z-A)":
        return df.sort_values(['Workbook Name', 'File Path'], ascending=[False, True])
    if sort_by == "Complexity (High to Low)":
        order = {'High': 0, 'Medium': 1, 'Low': 2, 'Manual Review': 3}
        tmp = df.copy()
        tmp['_complexity_rank'] = tmp['Complexity'].map(order).fillna(4)
        return tmp.sort_values(['_complexity_rank', 'Workbook Name'])
    if sort_by == "Duplicate Type":
        return df.sort_values(['Duplicate Type', 'Workbook Name'])
    if sort_by == "Relationship Count (High to Low)":
        return df.sort_values(['Relationships', 'Workbook Name'], ascending=[False, True])
    if sort_by == "Formula Count (High to Low)":
        return df.sort_values(['Calculated Fields', 'Workbook Name'], ascending=[False, True])
    return df.sort_values(['Workbook Name', 'File Path'])


if submit:
    st.session_state["analysis_has_run"] = True
    st.session_state.pop("workbook_explorer_choice", None)
    try:
        if mode == "Individual Workbook Files":
            if not uploaded_twb_files:
                st.error("Please upload at least one .twb or .twbx file.")
                st.stop()
            all_items = []
            for f in uploaded_twb_files:
                data = f.read()
                text = load_twb_text(data, f.name)
                all_items.append(("uploaded_files", f.name, text))
        else:
            if not uploaded_zip:
                st.error("Please upload a ZIP file.")
                st.stop()
            all_items = collect_twbx_from_zip(uploaded_zip.read())
            if not all_items:
                st.error("No .twb/.twbx files found inside the ZIP.")
                st.stop()

        usage_uploaded = usage_file is not None
        if usage_uploaded:
            allowed_usage_names = read_usage_file(usage_file)
        else:
            allowed_usage_names = set()

        if allowed_usage_names:
            selected_items, matched_pairs = filter_workbooks_by_usage(all_items, allowed_usage_names)
            if not selected_items:
                st.warning(
                    "Content usage names were detected, but none matched the uploaded workbooks. "
                    "KashMap will process all uploaded workbooks instead."
                )
                selected_items = all_items
        else:
            if usage_uploaded:
                st.warning("No workbook names could be detected from the content usage file. KashMap will process all uploaded workbooks.")
            else:
                st.info("No content usage file uploaded. This step is optional, so KashMap will process all uploaded workbooks.")
            selected_items = all_items
            allowed_usage_names = {normalize_workbook_name(name) for _, name, _ in all_items}
            matched_pairs = [(normalize_workbook_name(name), name) for _, name, _ in all_items]

        raw_text_lookup = {(folder, name): content for folder, name, content in selected_items}

        tds_map = {}
        if metadata_zip_file is not None:
            try:
                tds_map = collect_tds_metadata_from_zip(metadata_zip_file.read())
            except Exception as e:
                st.warning(f"Could not read TDS/TDSX Metadata ZIP: {e}")

        workbook_results = []
        errors = []
        for folder, name, text in selected_items:
            try:
                parsed = parse_twb(text)
            except Exception as e:
                parsed = None
                errors.append(f"{name}: {e}")
            workbook_results.append((folder, name, parsed))

        dup_records = build_workbook_duplicate_analysis(workbook_results)
        output_stream, display_tables, counts = build_output_workbook(
            workbook_results, dup_records, allowed_usage_names, matched_pairs, tds_map,
        )

        fn = output_name if output_name.lower().endswith(".xlsx") else f"{output_name}.xlsx"

        st.session_state["analysis_result"] = {
            "workbook_results": workbook_results, "dup_records": dup_records,
            "output_stream": output_stream, "display_tables": display_tables, "counts": counts,
            "errors": errors, "fn": fn, "raw_text_lookup": raw_text_lookup,
        }

    except Exception as e:
        st.error(str(e))

analysis_result = st.session_state.get("analysis_result")

if analysis_result:
    try:
        display_tables = analysis_result["display_tables"]
        overview_df = display_tables["overview"]
        fn = analysis_result["fn"]

        show_download_button(analysis_result["output_stream"], fn, "download_top")

        st.markdown('<div class="section-label">Choose a Tableau workbook to inspect</div>', unsafe_allow_html=True)
        if overview_df.empty:
            st.info("No workbook rows were available to inspect.")
        else:
            total_workbooks = len(overview_df)
            search_term, sort_choice = "", "Name (A-Z)"

            if total_workbooks >= 20:
                search_col, sort_col = st.columns([0.6, 0.4])
                with search_col:
                    search_term = st.text_input("Search workbooks by name", placeholder="Type part of a workbook name...", key="workbook_search_term")
                with sort_col:
                    sort_choice = st.selectbox(
                        "Sort by",
                        ["Name (A-Z)", "Name (Z-A)", "Complexity (High to Low)", "Duplicate Type",
                         "Relationship Count (High to Low)", "Formula Count (High to Low)"],
                        key="workbook_sort_choice",
                    )

            filtered_df = overview_df
            if search_term.strip():
                filtered_df = filtered_df[filtered_df['Workbook Name'].astype(str).str.contains(search_term.strip(), case=False, na=False)]
            if filtered_df.empty:
                st.warning(f"No workbooks matched '{search_term}'. Showing all {total_workbooks} workbooks instead.")
                filtered_df = overview_df
            filtered_df = sort_overview_df(filtered_df, sort_choice)

            workbook_options = [f"{row['Workbook Name']} | {row['File Path']} | {row['Duplicate Type']}" for _, row in filtered_df.iterrows()]
            if total_workbooks >= 20:
                st.caption(f"Showing {len(workbook_options)} of {total_workbooks} workbooks")

            selected_option = st.selectbox("Choose a workbook to inspect", workbook_options, key="workbook_explorer_choice", label_visibility="collapsed")
            selected_name, selected_path, _ = [p.strip() for p in selected_option.split('|', 2)]

            selected_overview = overview_df[(overview_df['Workbook Name'] == selected_name) & (overview_df['File Path'] == selected_path)]

            if not selected_overview.empty:
                workbook_row = selected_overview.iloc[0]
                parsed = None
                for folder, name, p in analysis_result["workbook_results"]:
                    if name == selected_name and folder == selected_path:
                        parsed = p
                        break

                st.subheader(f"Workbook Inspector: {selected_name}")

                if parsed is None:
                    st.warning("This workbook could not be parsed.")
                else:
                    r1, r2, r3, r4, r5 = st.columns(5)
                    with r1:
                        with st.container(border=True):
                            st.metric("Tables", sum(len(ds['tables']) for ds in parsed['datasources']))
                    with r2:
                        with st.container(border=True):
                            st.metric("Relationships", len(parsed['relationships']))
                    with r3:
                        with st.container(border=True):
                            st.metric("Calc Fields", len(parsed['calculated_fields']))
                    with r4:
                        with st.container(border=True):
                            st.metric("Filters", len(parsed['filters']))
                    with r5:
                        with st.container(border=True):
                            st.metric("Worksheets", len(parsed['worksheets']))

                    tables_df = build_tables_df(selected_path, selected_name, parsed)
                    rels_df = build_relationships_df(selected_path, selected_name, parsed)
                    formulas_df = build_formulas_df(selected_path, selected_name, parsed)
                    lods_df = build_lod_df(selected_path, selected_name, parsed)
                    filters_df = build_filters_df(selected_path, selected_name, parsed)
                    worksheets_df = build_worksheets_df(selected_path, selected_name, parsed)
                    dashboards_df = build_dashboards_df(selected_path, selected_name, parsed)
                    actions_df = build_actions_df(selected_path, selected_name, parsed)

                    single_wb_excel = build_single_workbook_excel(
                        workbook_row, tables_df, rels_df, formulas_df, lods_df,
                        filters_df, worksheets_df, dashboards_df, actions_df,
                    )
                    download_key = re.sub(r'[^A-Za-z0-9_]+', '_', f"{selected_name}_{selected_path}")
                    file_name = re.sub(r'[^A-Za-z0-9_.-]+', '_', Path(str(selected_name)).stem) + "_workbook.xlsx"
                    st.download_button(
                        "Download This Workbook (Excel)", data=single_wb_excel, file_name=file_name,
                        mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                        key=f"download_single_wb_{download_key}", use_container_width=True,
                    )

                    workbook_tabs = st.tabs([
                        "Start Here", "Sources & Relationships", "Filters", "Calculations & LOD",
                        "Worksheets & Dashboards", "Power BI Build Plan", "Table → Database Mapping",
                    ])

                    with workbook_tabs[0]:
                        render_start_here(workbook_row, parsed)
                        render_duplicate_code_diff(
                            workbook_row, display_tables['duplicates'], overview_df,
                            analysis_result.get('raw_text_lookup', {}), selected_name, selected_path,
                        )

                    with workbook_tabs[1]:
                        st.subheader("Tables")
                        show_table_or_info(tables_df, "No source tables detected.")
                        st.subheader("Relationships / Joins")
                        show_table_or_info(rels_df, "No relationships/joins detected.")
                        if not rels_df.empty:
                            st.warning("Tableau relationships do not declare cardinality - confirm key uniqueness before building the Power BI model relationship.")

                    with workbook_tabs[2]:
                        show_table_or_info(filters_df, "No filters detected.")

                    with workbook_tabs[3]:
                        st.subheader("Calculated Fields")
                        show_table_or_info(formulas_df, "No calculated fields detected.")
                        if not lods_df.empty:
                            st.subheader("LOD Expression Translations (verify caveats before use)")
                            show_table_or_info(lods_df, "No LOD expressions detected.")

                    with workbook_tabs[4]:
                        st.subheader("Worksheets")
                        show_table_or_info(worksheets_df, "No worksheets detected.")
                        st.subheader("Dashboards")
                        show_table_or_info(dashboards_df, "No dashboards detected.")
                        st.subheader("Actions")
                        show_table_or_info(actions_df, "No actions detected.")

                    with workbook_tabs[5]:
                        st.info("The full evidence workbook is available from the Download Full Excel Report button.")
                        ai_tab, m_tab, validation_tab, review_tab = st.tabs([
                            "AI Build Plan", "Power Query M Draft", "Validation", "Manual Review",
                        ])
                        with ai_tab:
                            render_ai_build_plan(workbook_row, parsed)
                        with m_tab:
                            st.markdown(
                                "Tableau has no HOLD-chain equivalent, so instead of SQL CTEs this is a rule-based "
                                "**Power Query M merge draft** built from the workbook's relationships."
                            )
                            m_draft = generate_power_query_m_draft(workbook_row, parsed)
                            st.code(m_draft, language='text')
                            st.download_button(
                                "Download .m draft", data=m_draft.encode('utf-8'),
                                file_name=f"{re.sub(r'[^A-Za-z0-9_]+', '_', selected_name)}_draft.m",
                                mime="text/plain", key=f"download_m_{download_key}",
                            )
                        with validation_tab:
                            # Validation rows are only keyed by workbook NAME (build_validation_rows_raw
                            # never receives a folder/path), so if two different uploaded workbooks happen
                            # to share the same file name in different folders, this filters by name only -
                            # both would show the same checklist. Rare edge case, flagged here rather than
                            # silently mismatching.
                            validation_df = display_tables['validation'][
                                display_tables['validation']['Workbook Name'] == selected_name
                            ]
                            st.subheader("Migration Validation Checklist")
                            show_table_or_info(
                                validation_df[['Check Category', 'Validation Check', 'Trigger Rule']] if not validation_df.empty else validation_df,
                                "No validation checks generated for this workbook.",
                            )
                        with review_tab:
                            render_manual_review(parsed)

                    with workbook_tabs[6]:
                        table_db_df = display_tables.get('table_db_mapping')
                        if table_db_df is not None and not table_db_df.empty:
                            st.caption("This table applies to all workbooks - it's the same regardless of which workbook is selected above.")
                            m1, m2, m3 = st.columns(3)
                            m1.metric("Total Unique Tables", len(table_db_df))
                            m2.metric("Mapped", display_tables.get("table_db_mapped_count", 0))
                            m3.metric("Needs Review", display_tables.get("table_db_needs_review_count", 0))
                            show_table_or_info(table_db_df, "No table mapping rows found.", large=True)
                            mapping_csv = table_db_df.to_csv(index=False).encode('utf-8')
                            st.download_button("Download Table DB Mapping CSV", data=mapping_csv, file_name="table_db_mapping.csv", mime="text/csv", key="download_table_db_mapping")
                        else:
                            st.info("No TDS/TDSX Metadata ZIP was uploaded, so no table-to-database mapping is available. Upload a TDS/TDSX ZIP and re-analyze to see this data.")

        if analysis_result["errors"]:
            st.warning(f"{len(analysis_result['errors'])} file(s) had errors.")
            with st.expander("View Error Log"):
                for err in analysis_result["errors"]:
                    st.text(err)

    except Exception as e:
        st.error(str(e))
